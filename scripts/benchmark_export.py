"""Offline synthetic export benchmark. No network or production DB/output writes."""
import argparse
from copy import deepcopy
import csv
import hashlib
from io import BytesIO
import json
import os
import shutil
from pathlib import Path
import tempfile
import time
import zipfile

from PIL import Image
from openpyxl import load_workbook
from bise_results.parser import parse_result
from bise_results.reports import export_batch, master_name
from bise_results.store import locked_store, import_batch
from bise_results.publishing import ArtifactDigests


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--rows',type=int,default=10000)
    parser.add_argument('--card-kib',type=int,default=320)
    parser.add_argument('--cold',action='store_true',help='Independent export without the collector receipt cache')
    parser.add_argument('--output',type=Path,default=Path('docs/evidence/export-performance.json'))
    args=parser.parse_args()
    if args.rows<1 or args.rows>100000 or args.card_kib<0:raise ValueError('Invalid benchmark dimensions')
    fixture=Path(__file__).resolve().parents[1]/'tests/fixtures/ssc-part1.html'
    base_html=fixture.read_text();base=parse_result(base_html,'001234',exam_type='1')
    image=BytesIO();Image.new('RGB',(20,20),'white').save(image,format='PNG')
    # Valid PNG plus incompressible trailing bytes models an already-compressed
    # card's IO size without using actual student images or network requests.
    card=image.getvalue()+os.urandom(max(0,args.card_kib*1024-len(image.getvalue())))
    card_sha=hashlib.sha256(card).hexdigest()
    scratch=Path(__file__).resolve().parents[1]/'artifacts/benchmarks'
    scratch.mkdir(parents=True,exist_ok=True)
    if shutil.disk_usage(scratch).free < args.rows*len(card)*4 + 256*1024*1024:
        raise ValueError('Not enough scratch disk space for this benchmark size')
    with tempfile.TemporaryDirectory(prefix='bise-export-benchmark-',dir=scratch) as directory:
        root=Path(directory);source=root/'rolls.csv'
        source.write_text('Roll Numbers\n'+'\n'.join(f'{i+1:06d}' for i in range(args.rows))+'\n')
        print(f'Preparing {args.rows:,} synthetic receipts ({len(card):,} bytes/card)...',flush=True)
        with locked_store(root/'results.sqlite3') as db:
            batch=import_batch(db,source,'Synthetic benchmark',exam_type='1')
            updates=[]
            receipts=ArtifactDigests()
            for i in range(args.rows):
                roll=f'{i+1:06d}';folder=root/'evidence'/roll;folder.mkdir(parents=True)
                html=base_html.replace('001234',roll);parsed=deepcopy(base)
                parsed.update(roll=roll,institution=f'Synthetic School {i%10:02d}')
                (folder/'card.html').write_text(html);(folder/'card.png').write_bytes(card)
                (folder/'parsed.json').write_text(json.dumps(parsed))
                receipt={name:{'path':str(folder/name),'sha256':card_sha if name=='card.png' else hashlib.sha256((folder/name).read_bytes()).hexdigest()} for name in ('card.html','card.png','parsed.json')}
                for name,meta in receipt.items():receipts.record_written(folder/name,meta['sha256'])
                updates.append((json.dumps(parsed),json.dumps(receipt),roll))
            with db:db.executemany("UPDATE jobs SET state='verified',result_json=?,receipt_json=? WHERE roll=?",updates)
            del updates
            began=time.perf_counter()
            destination,summary=export_batch(db,batch,root/'output',include_pending=False,progress=lambda x:print(x,flush=True),digests=None if args.cold else receipts)
            first=time.perf_counter()-began
            # Independent file/schema/link checks outside the timed export.
            manifest=json.loads((destination/'Data/manifest.json').read_text())
            with (destination/(master_name(2026,'1')+'.csv')).open(encoding='utf-8-sig',newline='') as stream:
                assert len(list(csv.DictReader(stream)))==args.rows
            book=load_workbook(destination/manifest['master_workbook'],read_only=True,data_only=True)
            assert book['All Students'].max_row==args.rows+1
            assert book['Subject Marks'].max_row==8*args.rows+1
            book.close()
            with zipfile.ZipFile(destination/'All Result Cards.zip') as archive:
                assert len(archive.namelist())==args.rows
                assert hashlib.sha256(archive.read(archive.namelist()[-1])).hexdigest()==card_sha
            files=[p for p in destination.rglob('*') if p.is_file()]
            stats=[p.stat() for p in files]
            physical={ (st.st_dev,st.st_ino):st.st_size for st in stats }
            before=(destination/manifest['master_workbook']).stat().st_mtime_ns
            began=time.perf_counter();export_batch(db,batch,root/'output',include_pending=False)
            rerun=time.perf_counter()-began
            assert (destination/manifest['master_workbook']).stat().st_mtime_ns==before
            report={'scenario':'independent_export' if args.cold else 'collector_stop','xlsx_backend':'xlsxwriter','synthetic_rows':args.rows,'subjects_per_student':8,'institutions':10,'card_bytes':len(card),
                    'first_full_export_seconds':round(first,3),'unchanged_rerun_seconds':round(rerun,3),
                    'stage_timings':manifest['export_timings'],'all_export_formats_and_links_preserved':True,
                    'unchanged_rerun_preserved_mtime':True,'logical_report_bytes':sum(st.st_size for st in stats),
                    'physical_report_file_bytes':sum(physical.values()),'network_requests':0,'production_database_modified':False}
            args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,indent=2)+'\n')
            print(json.dumps(report,indent=2),flush=True)


if __name__=='__main__':main()
