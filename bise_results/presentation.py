"""Readable Excel workbooks with overview and institution navigation."""

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .publishing import unique_names

LABELS = {
    'roll': 'Roll No', 'institution': 'Institution / District', 'name': 'Student Name',
    'father_name': 'Father Name', 'date_of_birth': 'Date of Birth',
    'student_cnic': 'Student CNIC / B-Form', 'father_cnic': 'Father CNIC',
    'registration_number': 'Registration No', 'group': 'Group',
    'obtained': 'Obtained Marks', 'maximum': 'Maximum Marks', 'percentage': 'Percentage',
    'result_statement': 'Board Result / Remarks', 'total_raw': 'Board Total Details',
    'declared_on': 'Result Declared On', 'fetched_at': 'Retrieved At',
    'result_status': 'Result Status', 'data_completeness': 'Available Details',
    'data_notes': 'Additional Notes', 'state': 'Collection Status', 'error': 'Lookup Issue',
    'result_card': 'Result Card',
}
SUBJECT_FIELDS = ['subject', 'max_theory_raw', 'max_practical_raw', 'max_total_raw', 'theory_i_raw', 'theory_ii_raw', 'practical_ii_raw', 'obtained_raw', 'percentile_raw', 'grade', 'status']
SUBJECT_LABELS = ['Subject', 'Maximum Theory', 'Maximum Practical', 'Maximum Total', 'Theory Part I', 'Theory Part II', 'Practical Part II', 'Obtained Total', 'Percentile Score', 'Relative Grade', 'Subject Status']
STATE_LABELS = {'verified': 'Saved', 'pending': 'Pending', 'running': 'In progress', 'needs_review': 'Identity unavailable', 'failed': 'Lookup failed'}
DETAIL_LABELS = {'complete': 'Full marks available', 'partial': 'Some marks unavailable', 'summary_only': 'Board status only', 'identity_only': 'Student details only'}
GREEN = '1C4D24'


def display(record, column):
    if column == 'result_card':
        return f"Result Cards/{record['roll']}.png" if record['state'] == 'verified' and not record.get('_database_only') else ''
    value = record.get(column)
    if column == 'state':
        return STATE_LABELS.get(value, value)
    if column == 'data_completeness':
        return DETAIL_LABELS.get(value, value)
    return value


def style_table(sheet, table_name):
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    # Widths depend on headers only. Iterating sheet.columns touches every cell.
    for cell in sheet[1]:
        heading = cell.value
        width = {'Roll No': 12, 'Institution / District': 28, 'Student Name': 27, 'Father Name': 29,
                 'Date of Birth': 16, 'Student CNIC / B-Form': 25, 'Father CNIC': 23,
                 'Registration No': 25, 'Board Result / Remarks': 60,
                 'Board Total Details': 45, 'Additional Notes': 55, 'Retrieved At': 28}.get(heading, 22)
        sheet.column_dimensions[cell.column_letter].width = width
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
        sheet.add_table(table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.outlinePr.summaryRight = False


def student_sheet(book, name, records, table_name, columns, safe_text, text_columns):
    columns = [column for column in columns if column != 'result_card']
    sheet = book.create_sheet(name)
    sheet.append([LABELS[column] for column in columns])
    for row_number, record in enumerate(records, start=2):
        sheet.append([safe_text(display(record, column)) for column in columns])
        for index, column in enumerate(columns, start=1):
            # Worksheet.max_row scans all existing cells: never call it per cell.
            cell = sheet.cell(row_number, index)
            if column in text_columns:
                cell.number_format = '@'
            if column == 'percentage':
                cell.number_format = '0.00"%"'
            if column == 'result_status' and cell.value:
                color = 'FCE4D6' if cell.value == 'FAIL' else 'FFF2CC' if cell.value in ('ABSENT', 'MARKS NOT IMPROVED') else 'E2EFDA'
                cell.fill = PatternFill('solid', fgColor=color)
    style_table(sheet, table_name)
    return sheet


def build_workbook(path, records, summary, title, columns, safe_text, text_columns, institutions=(), queue_summary=None):
    book = Workbook()
    overview = book.active
    overview.title = 'Overview'
    overview.sheet_view.showGridLines = False
    overview.merge_cells('A1:G1')
    overview['A1'] = title
    overview['A1'].font = Font(size=21, bold=True, color=GREEN)
    overview.row_dimensions[1].height = 38
    overview.merge_cells('A2:G2')
    overview['A2'] = 'Start here. Open All Students or an institution tab. Cards are in the Result Cards folder.'
    overview['A2'].font = Font(size=12, color='555555')
    overview.append([])
    queue_summary = queue_summary or summary
    overview.append(['Records in master', summary.get('total', 0), 'Saved', summary.get('verified', 0), 'Not yet saved', queue_summary.get('total', 0) - queue_summary.get('verified', 0)])
    overview.append(['All Students', 'Open full student list', 'Subject Marks', 'Open subject breakdown'])
    for cell, target in [('A5', 'All Students'), ('B5', 'All Students'), ('C5', 'Subject Marks'), ('D5', 'Subject Marks')]:
        overview[cell].hyperlink = f"#'{target}'!A1"
        overview[cell].style = 'Hyperlink'
    overview.append(['Imported rolls', queue_summary.get('total', 0)])
    overview.merge_cells('A7:G7')
    overview['A7'] = 'Student details are retained even when the board provides no numeric total.'
    overview.append([])
    overview.append(['Institution / District', 'Students', 'Numeric total available', 'Student tab', 'Separate workbook'])
    groups = defaultdict(list)
    for record in records:
        if record['state'] == 'verified':
            groups[record.get('institution', '')].append(record)
    sheet_names = unique_names(groups, fallback='Institution not provided', limit=31,
                               reserved=('Overview', 'All Students', 'Subject Marks', 'Missing Results'))
    paths = {entry['institution_district']: entry for entry in institutions}
    for row, (institution, members) in enumerate(sorted(groups.items()), start=10):
        tab = sheet_names[institution]
        entry = paths.get(institution)
        overview.append([safe_text(institution or 'Institution not provided'), len(members),
                         sum(record.get('obtained') is not None for record in members), 'Open students', 'Open report' if entry else ''])
        overview.cell(row, 4).hyperlink = f"#'{tab.replace(chr(39), chr(39)*2)}'!A1"
        overview.cell(row, 4).style = 'Hyperlink'
        if entry:
            overview.cell(row, 5).hyperlink = str(Path(entry['folder']) / entry['workbook'])
            overview.cell(row, 5).style = 'Hyperlink'
    overview.append([])
    overview.append(['Board outcome', 'Students'])
    for status, count in sorted(Counter(record.get('result_status') or 'Outcome unavailable' for record in records if record['state'] == 'verified').items()):
        overview.append([safe_text(status), count])
    overview.freeze_panes = 'A10'
    for letter, width in {'A': 35, 'B': 24, 'C': 29, 'D': 24, 'E': 26, 'F': 12, 'G': 12}.items():
        overview.column_dimensions[letter].width = width
    for row in (4, 9):
        for cell in overview[row]:
            cell.fill = PatternFill('solid', fgColor='E2EFDA')
            cell.font = Font(bold=True, color=GREEN)

    student_sheet(book, 'All Students', records, 'AllStudents', columns, safe_text, text_columns)
    for number, (institution, members) in enumerate(sorted(groups.items()), start=1):
        student_sheet(book, sheet_names[institution], members, f'Institution{number}', columns, safe_text, text_columns)
    subjects = book.create_sheet('Subject Marks')
    subjects.append(['Roll No', 'Institution / District', *SUBJECT_LABELS])
    subject_row = 1
    for record in records:
        for subject in record.get('subjects', []):
            subjects.append([record['roll'], safe_text(record.get('institution')), *[safe_text(subject.get(field)) for field in SUBJECT_FIELDS]])
            subject_row += 1
            subjects.cell(subject_row, 1).number_format = '@'
    style_table(subjects, 'SubjectMarks')
    missing = book.create_sheet('Missing Results')
    missing.append(['Roll No', 'Collection Status', 'Lookup Issue'])
    missing_row = 1
    for record in records:
        if record['state'] != 'verified':
            missing.append([record['roll'], display(record, 'state'), safe_text(record.get('error'))])
            missing_row += 1
            missing.cell(missing_row, 1).number_format = '@'
    style_table(missing, 'MissingResults')
    book.active = 0
    book.save(path)
