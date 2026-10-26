"""Performance contracts without fragile wall-clock assertions."""
from collections import Counter
from copy import deepcopy
import json
import os
from pathlib import Path
import tempfile
import unittest
import sys
from contextlib import redirect_stdout
from io import StringIO
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from bise_results.parser import parse_result
from bise_results.publishing import ArtifactDigests,file_hashes
from bise_results.reports import export_batch,workbook
from bise_results.store import import_batch,jobs_for,locked_store,check_receipt
from test_identity_fields import FIXTURE,receipt_for


class ExportPerformanceTests(unittest.TestCase):
    def test_workbook_does_not_scan_all_cells_per_record(self):
        base=parse_result(FIXTURE.read_text(),'001234')
        records=[dict(deepcopy(base),roll=f'{n:06d}',state='verified') for n in range(1,301)]
        original=Worksheet.max_row.fget;calls=0
        def counted(sheet):
            nonlocal calls
            calls+=1
            return original(sheet)
        with tempfile.TemporaryDirectory() as temp:
            target=Path(temp)/'report.xlsx'
            with patch.object(Worksheet,'max_row',property(counted)):
                workbook(target,records,{'verified':300,'total':300})
            # The old implementation made more than 15,000 full-sheet scans.
            self.assertLess(calls,100)
            book=load_workbook(target,read_only=True)
            self.assertEqual(book['All Students'].max_row,301)
            self.assertEqual(book['Subject Marks'].max_row,2401)
            book.close()

    def test_missing_result_rows_do_not_rescan_the_sheet(self):
        original=Worksheet.max_row.fget;calls=0
        def counted(sheet):
            nonlocal calls
            calls+=1
            return original(sheet)
        records=[{'roll':f'{i:06d}','state':'failed','error':'Synthetic failure'} for i in range(1,401)]
        with tempfile.TemporaryDirectory() as temp:
            with patch.object(Worksheet,'max_row',property(counted)):
                workbook(Path(temp)/'report.xlsx',records,{'failed':400,'total':400})
        self.assertLess(calls,100)

    def test_saved_data_can_be_exported_by_class_without_browser_or_batch_id(self):
        from bise_results.__main__ import main
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'rolls.csv';source.write_text('Roll Numbers\n001234\n')
            database=root/'state.sqlite3'
            with locked_store(database) as db:
                import_batch(db,source,'Part I',exam_type='1')
                import_batch(db,source,'Part II',exam_type='2')
            argv=['bise_results','export','--year','2026','--class','9','--database',str(database),'--output',str(root/'output')]
            with patch.object(sys,'argv',argv),redirect_stdout(StringIO()):self.assertEqual(main(),0)
            self.assertTrue((root/'output/Matric 2026 - Part-I Annual/Data/manifest.json').exists())
            self.assertFalse((root/'output/Matric 2026 - Part-II Annual').exists())

    def test_digest_cache_detects_same_size_mutation_even_when_mtime_is_restored(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'file.bin';path.write_bytes(b'original')
            digest=ArtifactDigests();first=digest(path);stamp=path.stat()
            self.assertEqual(digest(path),first);self.assertEqual(digest.reads,1)
            path.write_bytes(b'changed!');os.utime(path,ns=(stamp.st_atime_ns,stamp.st_mtime_ns))
            self.assertNotEqual(digest(path),first);self.assertEqual(digest.reads,2)

    def test_collector_hash_cache_does_not_hide_changed_evidence_at_stop(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'rolls.csv';source.write_text('Roll Numbers\n001234\n')
            with locked_store(root/'state.sqlite3') as db:
                batch=import_batch(db,source,'A');evidence=root/'evidence';evidence.mkdir()
                result=parse_result(FIXTURE.read_text(),'001234');receipt=receipt_for(evidence,result=result)
                with db:db.execute("update jobs set state='verified',result_json=?,receipt_json=?",(json.dumps(result),json.dumps(receipt)))
                cache=ArtifactDigests()
                for value in receipt.values():cache(value['path'])
                png=evidence/'card.png';old=png.stat();png.write_bytes(b'x'*old.st_size);os.utime(png,ns=(old.st_atime_ns,old.st_mtime_ns))
                destination,summary=export_batch(db,batch,root/'output',digests=cache)
                self.assertEqual(summary,{'needs_review':1,'total':1})
                self.assertFalse(json.loads((destination/'Data/manifest.json').read_text())['complete_collection'])

    def test_shared_batches_validate_original_artifact_bytes_only_once_per_export(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'rolls.csv';source.write_text('Roll Numbers\n001234\n')
            with locked_store(root/'state.sqlite3') as db:
                batch=import_batch(db,source,'A');import_batch(db,source,'B')
                evidence=root/'evidence';evidence.mkdir()
                result=parse_result(FIXTURE.read_text(),'001234');receipt=receipt_for(evidence,result=result)
                with db:db.execute("update jobs set state='verified',result_json=?,receipt_json=?",(json.dumps(result),json.dumps(receipt)))
                reads=Counter();original=ArtifactDigests.__call__
                def counted(cache,path):
                    before=cache.reads;value=original(cache,path)
                    if cache.reads>before:reads[str(Path(path).resolve())]+=1
                    return value
                log=[]
                with patch.object(ArtifactDigests,'__call__',counted):
                    destination,_=export_batch(db,batch,root/'output',progress=log.append)
                for value in receipt.values():self.assertEqual(reads[value['path']],1)
                self.assertTrue(any('writing master' in line for line in log))
                self.assertTrue(log[-1].startswith('Export complete:'))
                self.assertTrue(check_receipt(receipt))
                with zipfile.ZipFile(destination/'All Result Cards.zip') as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertEqual(archive.read('Result Cards/001234.png'),(evidence/'card.png').read_bytes())

    def test_generated_links_never_share_inodes_with_evidence_or_archived_versions(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'rolls.csv';source.write_text('Roll Numbers\n001234\n')
            with locked_store(root/'state.sqlite3') as db:
                batch=import_batch(db,source,'A');evidence=root/'evidence';evidence.mkdir()
                result=parse_result(FIXTURE.read_text(),'001234');receipt=receipt_for(evidence,result=result)
                with db:db.execute("update jobs set state='verified',result_json=?,receipt_json=?",(json.dumps(result),json.dumps(receipt)))
                destination,_=export_batch(db,batch,root/'output')
                card=destination/'Result Cards/001234.png';original=card.read_bytes()
                group=next((destination/'Institutions').iterdir())/'Result Cards/001234.png'
                self.assertNotEqual(card.stat().st_ino,(evidence/'card.png').stat().st_ino)
                self.assertEqual(card.stat().st_ino,group.stat().st_ino)
                result['name']='Updated Synthetic Name'
                with db:db.execute('update jobs set result_json=?',(json.dumps(result),))
                export_batch(db,batch,root/'output')
                archived=next((root/'artifacts/report-archive/Previous reports').glob('*/*/Files'))/'Result Cards/001234.png'
                self.assertNotEqual(card.stat().st_ino,archived.stat().st_ino)
                card.write_bytes(b'edited current generated card')
                self.assertEqual(archived.read_bytes(),original)
                self.assertTrue(check_receipt(receipt))

    def test_hardlink_failure_falls_back_to_independent_copies(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'rolls.csv';source.write_text('Roll Numbers\n001234\n')
            with locked_store(root/'state.sqlite3') as db:
                batch=import_batch(db,source,'A');evidence=root/'evidence';evidence.mkdir()
                result=parse_result(FIXTURE.read_text(),'001234');receipt=receipt_for(evidence,result=result)
                with db:db.execute("update jobs set state='verified',result_json=?,receipt_json=?",(json.dumps(result),json.dumps(receipt)))
                with patch('bise_results.reports.os.link',side_effect=OSError('no hardlinks')):
                    destination,_=export_batch(db,batch,root/'output')
                cards=list(destination.rglob('001234.png'))
                self.assertEqual(len(cards),2)
                self.assertEqual(cards[0].read_bytes(),cards[1].read_bytes())
                self.assertNotEqual(cards[0].stat().st_ino,cards[1].stat().st_ino)
