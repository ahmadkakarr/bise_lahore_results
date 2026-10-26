"""Interrupted collection must still refresh the master from committed rows."""

import asyncio
from argparse import Namespace
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook
from playwright.async_api import Error as BrowserError

from bise_results.parser import parse_result
from bise_results.reports import export_batch, master_name
from bise_results.runner import collect
from bise_results.store import import_batch, jobs_for, locked_store
from test_identity_fields import FIXTURE, receipt_for


class StopExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_ctrl_c_and_closed_driver_export_saved_rows_then_resume(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'rolls.csv'
            source.write_text('Roll Numbers\n001234\n001235\n001236\n')
            args = Namespace(csv=source, school='Stop test', year=2026, database=root/'results.sqlite3',
                             retry_unresolved=False, limit=0, workers=1, interval=3, captcha_mode='auto',
                             artifacts=root/'artifacts', output=root/'output', browser_executable=None, export_reports=True)
            context = SimpleNamespace(new_page=AsyncMock(return_value=MagicMock()), close=AsyncMock(side_effect=BrowserError('Connection closed while reading from the driver')))
            browser = SimpleNamespace(new_context=AsyncMock(return_value=context), close=AsyncMock(side_effect=BrowserError('Connection closed while reading from the driver')))
            manager = MagicMock()
            manager.__aenter__ = AsyncMock(return_value=SimpleNamespace(chromium=SimpleNamespace(launch=AsyncMock(return_value=browser))))
            manager.__aexit__ = AsyncMock(return_value=None)
            second_started = asyncio.Event()
            blocked = asyncio.Event()
            calls = []

            async def lookup(page, job, folder, *unused, **kwargs):
                calls.append(job['roll'])
                if job['roll'] == '001235' and not blocked.is_set():
                    second_started.set()
                    await blocked.wait()
                source_html = FIXTURE.read_text().replace('001234', job['roll'])
                parsed = parse_result(source_html, job['roll'])
                return parsed, receipt_for(folder, source_html, parsed)

            with patch('bise_results.runner.async_playwright', return_value=manager), patch('bise_results.runner.lookup', side_effect=lookup), redirect_stdout(io.StringIO()):
                task = asyncio.create_task(collect(args))
                await asyncio.wait_for(second_started.wait(), 1)
                task.cancel()  # asyncio.run receives this cancellation on first Ctrl+C.
                self.assertEqual(await task, 130)
                master = args.output/'Matric 2026 - Part-II Annual'/f'{master_name(2026)}.xlsx'
                self.assertEqual(load_workbook(master)['All Students'].max_row, 2)
                with locked_store(args.database) as db:
                    rows = db.execute('SELECT * FROM jobs ORDER BY roll').fetchall()
                    self.assertEqual([row['state'] for row in rows], ['verified', 'pending', 'pending'])
                    self.assertEqual(db.execute("SELECT count(*) FROM attempts WHERE outcome='interrupted'").fetchone()[0], 1)
                blocked.set()
                self.assertEqual(await collect(args), 0)
            self.assertEqual(calls.count('001234'), 1)
            self.assertEqual(load_workbook(master)['All Students'].max_row, 4)

    async def test_worker_start_failure_still_exports_prior_saved_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root/'rolls.csv'
            source.write_text('Roll Numbers\n001234\n001235\n')
            database = root/'results.sqlite3'
            with locked_store(database) as db:
                batch = import_batch(db, source, 'Failure test')
                folder = root/'saved'; folder.mkdir()
                parsed = parse_result(FIXTURE.read_text(), '001234')
                receipt = receipt_for(folder, result=parsed)
                with db:
                    db.execute("UPDATE jobs SET state='verified',result_json=?,receipt_json=? WHERE roll='001234'", (json.dumps(parsed),json.dumps(receipt)))
            browser = SimpleNamespace(new_context=AsyncMock(side_effect=RuntimeError('cannot create context')), close=AsyncMock())
            manager = MagicMock()
            manager.__aenter__ = AsyncMock(return_value=SimpleNamespace(chromium=SimpleNamespace(launch=AsyncMock(return_value=browser))))
            manager.__aexit__ = AsyncMock(return_value=None)
            args = Namespace(csv=source, school='Failure test', year=2026, database=database, retry_unresolved=False,
                             limit=0, workers=1, interval=3, captcha_mode='auto', artifacts=root/'artifacts',output=root/'output',browser_executable=None, export_reports=True)
            with patch('bise_results.runner.async_playwright', return_value=manager), redirect_stdout(io.StringIO()):
                with self.assertRaises(RuntimeError):
                    await collect(args)
            destination=root/'output/Matric 2026 - Part-II Annual'
            manifest=json.loads((destination/'Data/manifest.json').read_text())
            self.assertEqual(manifest['summary'], {'verified':1,'total':1})
            self.assertEqual(manifest['queue_summary']['pending'],1)
            self.assertFalse(manifest['complete_collection'])


class ProcessedExportTests(unittest.TestCase):
    def test_large_pending_queue_stays_out_of_master_but_counts_remain_visible(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            source=root/'rolls.csv'
            source.write_text('Roll Numbers\n'+'\n'.join(f'{roll:06d}' for roll in range(1000,2001))+'\n')
            with locked_store(root/'results.sqlite3') as db:
                batch=import_batch(db,source,'Large pending test')
                folder=root/'saved';folder.mkdir()
                html=FIXTURE.read_text().replace('001234','001000')
                parsed=parse_result(html,'001000');receipt=receipt_for(folder,html,parsed)
                with db:
                    db.execute("UPDATE jobs SET state='verified',result_json=?,receipt_json=? WHERE roll='001000'",(json.dumps(parsed),json.dumps(receipt)))
                destination,summary=export_batch(db,batch,root/'output',include_pending=False)
                self.assertEqual(summary,{'verified':1,'total':1})
                self.assertEqual(db.execute('SELECT count(*) FROM jobs').fetchone()[0],1001)
            manifest=json.loads((destination/'Data/manifest.json').read_text())
            self.assertEqual(manifest['queue_summary'],{'pending':1000,'verified':1,'total':1001})
            self.assertFalse(manifest['complete_collection'])
            book=load_workbook(destination/f'{master_name(2026)}.xlsx')
            self.assertEqual(book['All Students'].max_row,2)
            self.assertEqual(book['Overview']['B6'].value,1001)
            self.assertEqual(book['Overview']['F4'].value,1000)


if __name__=='__main__':
    unittest.main()
