"""Regressions for board identity fields and existing institution exports."""

import csv
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from PIL import Image

from bise_results.parser import parse_result
from bise_results.reports import write_group
from bise_results.presentation import LABELS
FIELD_NAMES = {label: field for field, label in LABELS.items()}

FIXTURE = Path(__file__).parent / "fixtures/ssc-part2.html"
IDENTITY = {
    "date_of_birth": "09/05/2008",
    "student_cnic": "00000-0123456-0",
    "father_cnic": "00000-0765432-0",
    "registration_number": "00001-00001-2020",
}


def receipt_for(root, source=None, result=None):
    source = source if source is not None else FIXTURE.read_text()
    result = result if result is not None else parse_result(source, "001234")
    (root / "card.html").write_text(source)
    Image.new("RGB", (2, 2), "white").save(root / "card.png")
    (root / "parsed.json").write_text(json.dumps(result))
    return {name: {"path": str(root / name), "sha256": hashlib.sha256((root / name).read_bytes()).hexdigest()} for name in ("card.html", "card.png", "parsed.json")}


class IdentityFieldTests(unittest.TestCase):
    def test_all_identity_fields_are_read_from_the_board(self):
        result = parse_result(FIXTURE.read_text(), "001234")
        for key, value in IDENTITY.items():
            with self.subTest(field=key):
                self.assertEqual(result.get(key), value)
                self.assertIsInstance(result[key], str)

    def test_missing_and_blank_fields_are_not_invented_or_shifted(self):
        source = FIXTURE.read_text().replace('<label id="lblDOB">09/05/2008</label>', "")
        source = source.replace("00000-0123456-0", "&nbsp;")
        result = parse_result(source, "001234")
        self.assertEqual(result.get("date_of_birth"), "")
        self.assertEqual(result.get("student_cnic"), "")
        self.assertEqual(result.get("father_cnic"), IDENTITY["father_cnic"])
        self.assertEqual(result["obtained"], 551)

    def test_nested_or_masked_values_keep_the_board_representation(self):
        source = FIXTURE.read_text().replace("00000-0123456-0", "<span>00000-*******-0</span>")
        result = parse_result(source, "001234")
        self.assertEqual(result.get("student_cnic"), "00000-*******-0")
        self.assertEqual(result.get("date_of_birth"), IDENTITY["date_of_birth"])

    def test_csv_excel_json_preserve_identity_fields_as_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            result = parse_result(FIXTURE.read_text(), "001234")
            result.update({"state": "verified", "error": None, "_receipt": receipt_for(root, result=result)})
            destination = root / "reports"
            write_group(destination, [result])
            with (destination / "results.csv").open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                csv_record = {FIELD_NAMES[key]: value for key, value in next(reader).items()}
                csv_headers = [FIELD_NAMES[key] for key in reader.fieldnames]
            book = load_workbook(destination / "results.xlsx", data_only=False)
            sheet = book["All Students"]
            columns = {FIELD_NAMES[cell.value]: cell.column for cell in sheet[1]}
            self.assertEqual([key for key in csv_headers if key != 'result_card'], list(columns))
            self.assertNotIn('result_card', columns)
            self.assertEqual(csv_record['result_card'], 'Result Cards/001234.png')
            json_record = json.loads((destination / "Data/results.json").read_text())[0]
            for key, value in IDENTITY.items():
                with self.subTest(field=key):
                    self.assertEqual(csv_record.get(key), value)
                    self.assertEqual(json_record.get(key), value)
                    cell = sheet.cell(2, columns[key])
                    self.assertEqual(cell.value, value)
                    self.assertEqual(cell.data_type, "s")
                    self.assertEqual(cell.number_format, "@")
            self.assertEqual(sheet.cell(2, columns["obtained"]).value, 551)
            self.assertEqual(sheet.cell(2, columns["obtained"]).data_type, "n")
            self.assertEqual(sheet.cell(2, columns["percentage"]).data_type, "n")
            self.assertEqual(book["Subject Marks"].max_row, 9)
            self.assertTrue((destination / "Result Cards.zip").is_file())

    def test_unprocessed_rows_have_blank_identity_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            write_group(root, [{"roll": "001234", "state": "pending"}], include_cards=False)
            with (root / "results.csv").open(encoding="utf-8-sig", newline="") as stream:
                row = {FIELD_NAMES[key]: value for key, value in next(csv.DictReader(stream)).items()}
            for key in IDENTITY:
                self.assertIn(key, row)
                self.assertEqual(row[key], "")
            self.assertEqual(row["roll"], "001234")


if __name__ == "__main__":
    unittest.main()
