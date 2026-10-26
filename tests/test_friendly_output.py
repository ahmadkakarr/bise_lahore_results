"""Friendly report navigation, stable publication and lossless migration."""

import csv
from html import escape
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from urllib.parse import unquote

from openpyxl import load_workbook

from bise_results.parser import parse_result
from bise_results.presentation import LABELS
from bise_results.publishing import archive_legacy, file_hashes, unique_names
from bise_results.reports import export_batch, master_name
from bise_results.store import import_batch, jobs_for, locked_store
from test_identity_fields import FIXTURE, IDENTITY, receipt_for


class FriendlyOutputTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.database = self.root / 'results.sqlite3'
        self.output = self.root / 'output'
        self.source = self.root / 'rolls.csv'
        self.source.write_text('Roll Numbers\n001234\n001235\n')
        with locked_store(self.database) as db:
            self.batch = import_batch(db, self.source, 'Input A')
            for row, institution in zip(jobs_for(db, self.batch), ('KASUR', 'LAHORE')):
                source = FIXTURE.read_text().replace('001234', row['roll']).replace('Example Institution / District', institution)
                result = parse_result(source, row['roll'])
                folder = self.root / f"card-{row['id']}"
                folder.mkdir()
                receipt = receipt_for(folder, source, result)
                with db:
                    db.execute("UPDATE jobs SET state='verified',result_json=?,receipt_json=? WHERE id=?", (json.dumps(result), json.dumps(receipt), row['id']))

    def export(self):
        with locked_store(self.database) as db:
            return export_batch(db, self.batch, self.output)

    def test_master_has_students_navigation_and_card_archives_without_excel_card_urls(self):
        destination, summary = self.export()
        self.assertEqual(summary, {'verified': 2, 'total': 2})
        self.assertEqual(destination.name, 'Matric 2026 - Part-II Annual')
        book = load_workbook(destination / f'{master_name(2026)}.xlsx')
        self.assertEqual(book.active.title, 'Overview')
        self.assertTrue({'All Students', 'KASUR', 'LAHORE', 'Subject Marks', 'Missing Results'}.issubset(book.sheetnames))
        self.assertEqual(book['All Students'].max_row, 3)
        self.assertEqual(book['KASUR'].max_row, 2)
        headers = {cell.value: cell.column for cell in book['All Students'][1]}
        for row in (2, 3):
            for key, value in IDENTITY.items():
                cell = book['All Students'].cell(row, headers[LABELS[key]])
                self.assertEqual(cell.value, value)
                self.assertEqual(cell.number_format, '@')
        self.assertNotIn('Result Card', headers)
        for name in ('All Students', 'KASUR', 'LAHORE'):
            self.assertTrue(all(cell.hyperlink is None for row in book[name] for cell in row))
        self.assertEqual(book['All Students'].freeze_panes, 'D2')
        self.assertEqual(len(book['All Students'].tables), 1)
        for cell in book['Overview']._cells.values():
            if cell.hyperlink is None:
                continue
            target = cell.hyperlink.target or ('#' + cell.hyperlink.location)
            if target.startswith('#'):
                self.assertIn(target.split("'")[1], book.sheetnames)
            else:
                self.assertTrue((destination / unquote(target).replace('\\','/')).is_file())
        for label in ('KASUR', 'LAHORE'):
            folder = destination / 'Institutions' / label
            report = load_workbook(folder / f'{label} Results.xlsx')
            self.assertNotIn('Result Card', [cell.value for cell in report['All Students'][1]])
            self.assertTrue(all(cell.hyperlink is None for row in report['All Students'] for cell in row))
            with zipfile.ZipFile(folder / f'{label} Result Cards.zip') as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(len(archive.namelist()), 1)
                self.assertTrue(archive.namelist()[0].startswith('Result Cards/'))

    def test_unchanged_rerun_reuses_files_without_creating_versions(self):
        destination, _ = self.export()
        original = file_hashes(destination)
        mtimes = {path: path.stat().st_mtime_ns for path in destination.rglob('*') if path.is_file()}
        again, _ = self.export()
        self.assertEqual(again, destination)
        self.assertEqual(file_hashes(again), original)
        self.assertTrue(all(path.stat().st_mtime_ns == stamp for path, stamp in mtimes.items()))
        self.assertEqual(list(self.output.iterdir()), [destination])
        self.assertFalse((self.root / 'artifacts/report-archive/Previous reports').exists())

    def test_previous_format_is_rebuilt_and_archived_after_card_url_removal(self):
        with patch('bise_results.reports.FORMAT_VERSION', 4):
            destination, _ = self.export()
        previous = json.loads((destination / 'Data/manifest.json').read_text())
        self.export()
        current = json.loads((destination / 'Data/manifest.json').read_text())
        self.assertEqual(current['format_version'], 5)
        self.assertNotEqual(previous['content_signature'], current['content_signature'])
        archives = list((self.root / 'artifacts/report-archive/Previous reports').glob('*/*/Files/Data/manifest.json'))
        self.assertEqual(len(archives), 1)
        self.assertEqual(json.loads(archives[0].read_text())['format_version'], 4)

    def test_full_board_institution_names_survive_safe_paths_without_batch_fallback(self):
        # Long, similar names collide when shortened for folders or Excel tabs.
        # The actual student field must retain the full board text in every report.
        prefix = 'Example Government Higher Secondary School / Girls & Boys: '
        for names in ((prefix + 'Central Campus ' * 4 + 'A',
                       prefix + 'Central Campus ' * 4 + 'B'), ('', 'KASUR')):
            with self.subTest(names=names):
                expected = {}
                with locked_store(self.database) as db:
                    for row, institution in zip(jobs_for(db, self.batch), names):
                        source = FIXTURE.read_text().replace('001234', row['roll']).replace(
                            'Example Institution / District', escape(institution))
                        result = parse_result(source, row['roll'])
                        self.assertEqual(result['institution'], institution)
                        receipt = receipt_for(self.root / f"card-{row['id']}", source, result)
                        with db:
                            db.execute('UPDATE jobs SET result_json=?,receipt_json=? WHERE id=?',
                                       (json.dumps(result), json.dumps(receipt), row['id']))
                        expected[row['roll']] = institution
                destination, _ = self.export()
                index = json.loads((destination / 'Data/institutions.json').read_text())
                self.assertEqual({entry['institution_district'] for entry in index}, set(names))
                self.assertEqual(len({entry['folder'] for entry in index}), 2)
                reports = [(destination, master_name(2026))]
                reports.extend((destination / entry['folder'], Path(entry['workbook']).stem)
                               for entry in index)
                for folder, report in reports:
                    with (folder / f'{report}.csv').open(encoding='utf-8-sig', newline='') as stream:
                        rows = list(csv.DictReader(stream))
                    for row in rows:
                        self.assertEqual(row[LABELS['institution']], expected[row['Roll No']])
                    book = load_workbook(folder / f'{report}.xlsx')
                    sheet = book['All Students']
                    columns = {cell.value: cell.column for cell in sheet[1]}
                    for row in range(2, sheet.max_row + 1):
                        roll = sheet.cell(row, columns['Roll No']).value
                        actual = sheet.cell(row, columns[LABELS['institution']]).value or ''
                        self.assertEqual(actual, expected[roll])
                    book.close()
                    for row in json.loads((folder / 'Data/results.json').read_text()):
                        self.assertEqual(row['institution'], expected[row['roll']])

    def test_changed_output_replaces_current_and_archives_previous_bytes(self):
        destination, _ = self.export()
        original = file_hashes(destination)
        with locked_store(self.database) as db:
            with db:
                db.execute("UPDATE jobs SET state='pending' WHERE roll='001235'")
        again, summary = self.export()
        self.assertEqual(again, destination)
        self.assertEqual(summary, {'verified': 1, 'pending': 1, 'total': 2})
        archived = list((self.root / 'artifacts/report-archive/Previous reports').glob('*/*/Files'))
        self.assertEqual(len(archived), 1)
        self.assertEqual(file_hashes(archived[0]), original)
        self.assertEqual(list(self.output.iterdir()), [destination])

    def test_publish_failure_restores_previous_report(self):
        destination, _ = self.export()
        original = file_hashes(destination)
        with locked_store(self.database) as db:
            with db:
                db.execute("UPDATE jobs SET state='pending' WHERE roll='001235'")
        with patch('bise_results.publishing.publish_staging', side_effect=OSError('simulated publish failure')):
            with self.assertRaises(OSError):
                self.export()
        self.assertEqual(file_hashes(destination), original)
        self.assertFalse(any(self.output.glob('.building-*')))

    def test_legacy_snapshots_are_archived_and_unrelated_files_preserved(self):
        legacy = self.output / self.batch / '20261020T123456123456'
        legacy.mkdir(parents=True)
        (legacy / 'manifest.json').write_text(json.dumps({'batch': self.batch}))
        (legacy / 'old-report.csv').write_text('historical data\n')
        original = file_hashes(legacy)
        unrelated = self.output / 'My Notes'
        unrelated.mkdir()
        (unrelated / 'note.txt').write_text('keep this')
        destination, _ = self.export()
        self.assertFalse((self.output / self.batch).exists())
        archived = list((self.root / 'artifacts/report-archive/Legacy output').glob('*/*/Files'))
        self.assertEqual(file_hashes(archived[0]), original)
        inventory = json.loads((archived[0].parent / 'Inventory.json').read_text())
        self.assertEqual(inventory['files'], original)
        self.assertEqual((unrelated / 'note.txt').read_text(), 'keep this')
        self.assertTrue(destination.is_dir())

    def test_multiple_lists_share_one_exam_master_without_duplicate_students(self):
        extra = self.root / 'more.csv'
        extra.write_text('Roll Numbers\n001234\n001236\n')
        with locked_store(self.database) as db:
            second = import_batch(db, extra, 'Input B')
            destination, summary = export_batch(db, second, self.output)
        self.assertEqual(summary['total'], 3)
        with (destination / f'{master_name(2026)}.csv').open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row['Roll No'] for row in rows], ['001234', '001235', '001236'])
        with locked_store(self.database) as db:
            again, _ = export_batch(db, self.batch, self.output)
        self.assertEqual(again, destination)
        self.assertEqual(len(list(self.output.iterdir())), 1)

    def test_safe_names_handle_collisions_reserved_names_and_blank_institutions(self):
        names = unique_names(['KASUR', 'LAHORE', 'A/B', 'A:B', 'a-b', '', '../', 'CON'])
        self.assertEqual(names['KASUR'], 'KASUR')
        self.assertEqual(names['LAHORE'], 'LAHORE')
        self.assertEqual(len({value.casefold() for value in names.values()}), len(names))
        self.assertTrue(all('/' not in value and ':' not in value and value not in ('.', '..') for value in names.values()))

    def test_unmanaged_destination_is_not_overwritten(self):
        destination = self.output / 'Matric 2026 - Part-II Annual'
        destination.mkdir(parents=True)
        note = destination / 'My File.txt'
        note.write_text('user data')
        with self.assertRaises(RuntimeError):
            self.export()
        self.assertEqual(note.read_text(), 'user data')


if __name__ == '__main__':
    unittest.main()
