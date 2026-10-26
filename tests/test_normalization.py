"""Offline normalization contracts, read-only inputs and repeatable publication."""

from copy import deepcopy
from datetime import datetime
import csv
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from bise_normalize import COLUMNS
from bise_normalize.app import execute, inspect_source
from bise_normalize.config import initialize, load, template, validate
from bise_normalize.engine import born_on, clean_phone, identifier
from bise_normalize.publish import hashes, locked
from bise_normalize.sources import capture, digest

SOURCE_HEADERS = ['Roll No', 'Institution / District', 'Student Name', 'Father Name',
                  'Date of Birth', 'Student CNIC / B-Form', 'Father CNIC',
                  'Collection Status', 'Result Status', 'Obtained Marks']


def student(roll='001234', institution='LAHORE', bform='00000-0123456-1', **extra):
    row = dict(zip(SOURCE_HEADERS, [roll, institution, 'Example Student', 'Example Father',
                   '09/05/2008', bform, '00000-0765432-1', 'Saved', 'ABSENT', '']))
    row.update(extra)
    return row


class NormalizationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'master.csv'
        self.sample = self.root / 'sample.csv'
        self.output = self.root / 'normalized_output/Lahore'
        self.profile = self.root / 'profiles/lahore.json'
        self.write_csv(self.sample, COLUMNS, [])
        self.rows = [student(), student('001235', bform='00000-0123457-2', **{'Date of Birth': '21/10/2010'}),
                     student('001236', 'KASUR'), student('001237', **{'Collection Status': 'Lookup failed', 'Student Name': ''})]
        self.write_source()
        initialize(self.profile, 'Lahore', self.source, self.sample, self.output)
        self.config = load(self.profile)
        self.config['output']['write_excel'] = True
        self.config['selection']['institutions'] = ['LAHORE']
        self.config['defaults'].update(Category='Schools (Secondary)', SubCategory='Open Category',
            Division='Lahore', District='Lahore', Tehsil='Shalimar', UC='UC No. 178, Manawan', Sport='Athletics')
        self.save_config()

    def write_csv(self, path, fields, rows):
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    def write_source(self):
        fields = list(dict.fromkeys([*SOURCE_HEADERS, *(k for row in self.rows for k in row)]))
        self.write_csv(self.source, fields, self.rows)

    def save_config(self):
        self.profile.write_text(json.dumps(self.config, indent=2))

    def run_normalize(self, mode='run'):
        self.save_config()
        return execute(self.profile, mode)

    def exported(self, file='Normalized.csv'):
        with (self.output / file).open(encoding='utf-8-sig', newline='') as stream:
            return list(csv.DictReader(stream))

    def provenance(self):
        return [json.loads(line) for line in (self.output / 'Data/provenance.jsonl').read_text().splitlines()]

    def test_lahore_export_exact_schema_defaults_origins_and_source_preserved(self):
        source_hash = digest(self.source)
        result = self.run_normalize()
        self.assertEqual(result['summary']['normalized'], 2)
        self.assertEqual(result['summary']['complete'], 2)
        rows = self.exported()
        self.assertEqual(list(rows[0]), COLUMNS)
        self.assertEqual(rows[0]['BForm'], '0000001234561')
        self.assertEqual(rows[0]['DOB'], '2008-05-09')
        self.assertEqual(rows[0]['Gender'], 'Male')
        self.assertEqual(rows[1]['Gender'], 'Female')
        self.assertEqual(rows[0]['Address'], 'LAHORE')
        self.assertEqual(rows[0]['Phone'], '03211234567')
        self.assertEqual(rows[0]['Institute'], '')
        self.assertEqual(rows[0]['SubCategory'], 'Open Category')
        self.assertEqual(rows[0]['JoggersSize'], '42')
        self.assertEqual(rows[1]['JoggersSize'], '38')
        self.assertEqual(digest(self.source), source_hash)
        info = self.provenance()[0]
        self.assertEqual(info['origins']['Gender'], 'student_bform_parity_unverified')
        self.assertEqual(info['origins']['Address'], 'board_institution_district_label')
        self.assertEqual(info['origins']['Phone'], 'operator_shared_contact')
        book = load_workbook(self.output / 'Normalized.xlsx')
        self.addCleanup(book.close)
        columns = {cell.value: cell.column for cell in book['Normalized'][1]}
        for field in ('Phone', 'BForm', 'JoggersSize'):
            cell = book['Normalized'].cell(2, columns[field])
            self.assertEqual(cell.data_type, 's')
            self.assertEqual(cell.number_format, '@')
        self.assertFalse((self.root / 'data/results.sqlite3').exists())

    def test_identical_run_and_json_key_reordering_leave_hashes_and_mtimes_unchanged(self):
        self.run_normalize()
        old = hashes(self.output)
        stamps = {p: p.stat().st_mtime_ns for p in self.output.rglob('*') if p.is_file()}
        self.profile.write_text(json.dumps(self.config, sort_keys=True, separators=(',', ':')))
        with patch('bise_normalize.app.prepare', side_effect=AssertionError('unchanged run must not transform')):
            result = execute(self.profile)
        self.assertEqual(result['status'], 'already_up_to_date')
        self.assertEqual(old, hashes(self.output))
        self.assertTrue(all(p.stat().st_mtime_ns == stamp for p, stamp in stamps.items()))

    def test_source_append_rebuilds_without_appending_duplicates_and_archives_bytes(self):
        self.run_normalize()
        old = hashes(self.output)
        self.rows.append(student('001238', bform='00000-0123458-1'))
        self.write_source()
        self.assertEqual(self.run_normalize()['summary']['normalized'], 3)
        archived = list((self.output.parent / '.normalization/previous').glob('*/*/Files'))
        self.assertEqual(len(archived), 1)
        self.assertEqual(hashes(archived[0]), old)
        self.assertEqual(execute(self.profile)['status'], 'already_up_to_date')

    def test_config_change_applies_and_missing_location_keeps_partial_rows(self):
        self.run_normalize()
        self.config['defaults']['UC'] = None
        result = self.run_normalize()
        self.assertEqual(result['summary']['complete'], 0)
        self.assertEqual(len(self.exported()), 2)
        self.assertEqual(len(self.exported('Complete Records.csv')), 0)
        book = load_workbook(self.output / 'Missing Details.xlsx')
        self.addCleanup(book.close)
        self.assertEqual(book.active.max_row, 3)
        self.assertIn('missing:UC', book.active.cell(2, 3).value)

    def test_size_boundaries_birthdays_unknown_gender_and_explicit_sizes(self):
        self.rows = [student('001234', **{'Date of Birth': '20/10/2011'}),
                     student('001235', bform='00000-0123457-2', **{'Date of Birth': '21/10/2011', 'Gender': 'Other'}),
                     student('001236', **{'KitSize': 'Small', 'JoggersSize': '37'})]
        self.write_source()
        self.run_normalize()
        rows = self.exported()
        self.assertEqual(rows[0]['JoggersSize'], '41')
        self.assertEqual(rows[1]['Gender'], 'Other')
        self.assertEqual(rows[1]['JoggersSize'], '')
        self.assertEqual(rows[2]['JoggersSize'], '37')
        self.assertEqual(rows[2]['KitSize'], 'Small')

    def test_invalid_identity_values_keep_student_without_using_father_cnic(self):
        self.rows = [student(bform='00000-*******-2', **{'Date of Birth': '31/02/2008'})]
        self.write_source()
        self.run_normalize()
        row = self.exported()[0]
        self.assertEqual(row['FullName'], 'Example Student')
        self.assertEqual(row['BForm'], '')
        self.assertEqual(row['DOB'], '')
        self.assertEqual(row['Gender'], '')
        self.assertEqual(row['JoggersSize'], '')
        self.assertIn('invalid_bform', self.provenance()[0]['issues'])

    def test_explicit_gender_phone_address_and_school_take_priority(self):
        self.rows = [student(**{'Gender': 'Female', 'Phone': '+92 300 1234567', 'Address': 'Operator supplied address'})]
        self.write_source()
        full_name = 'Example Government Girls School / Campus & District ' * 3
        self.config['participation']['school_name_mappings'] = {'LAHORE': full_name}
        self.config['defaults']['SubCategory'] = 'Private'
        self.run_normalize()
        row = self.exported()[0]
        self.assertEqual(row['Gender'], 'Female')
        self.assertEqual(row['Phone'], '03001234567')
        self.assertEqual(row['Address'], 'Operator supplied address')
        self.assertEqual(row['Institute'], full_name)
        self.assertEqual(row['SubCategory'], 'Private')
        self.assertIn('explicit_gender_conflicts_with_requested_parity', self.provenance()[0]['warnings'])

    def test_duplicate_bform_excludes_both_from_complete_without_losing_rows(self):
        self.rows = [student(), student('001235')]
        self.write_source()
        result = self.run_normalize()
        self.assertEqual(result['summary']['normalized'], 2)
        self.assertEqual(result['summary']['complete'], 0)
        self.assertEqual(result['summary']['duplicate_entries'], 2)
        self.assertEqual(len(self.exported()), 2)
        self.assertTrue(all('duplicate_target_bform' in r['issues'] for r in self.provenance()))

    def test_repeated_source_identity_collapses_and_conflicts_remain_visible(self):
        self.rows = [student(), student(), student(**{'Student Name': 'Different explicit name'})]
        self.write_source()
        self.run_normalize()
        self.assertEqual(len(self.exported()), 1)
        info = self.provenance()[0]
        self.assertEqual(info['source_rows'], [2, 3, 4])
        self.assertIn('conflicting_source_identity', info['issues'])
        self.assertEqual(info['source_variants'][0]['raw']['Student Name'], 'Different explicit name')

    def test_selection_ranges_are_logical_csv_records_and_roll_filters_compose(self):
        self.rows[0]['Student Name'] = 'Example\nStudent'
        self.write_source()
        self.config['selection']['row_ranges'] = [[3, 3]]
        self.config['selection']['roll_bounds'] = ['001235', '001235']
        self.config['selection']['result_statuses'] = ['ABSENT']
        self.run_normalize()
        self.assertEqual(self.provenance()[0]['roll'], '001235')
        self.assertEqual(len(self.exported()), 1)

    def test_xlsx_sheet_cell_range_and_date_cells(self):
        path = self.root / 'master.xlsx'
        book = Workbook()
        book.active.title = 'Overview'
        page = book.create_sheet('LAHORE')
        page.append(SOURCE_HEADERS)
        for index, row in enumerate(self.rows[:2]):
            values = [row[k] for k in SOURCE_HEADERS]
            if index == 0:
                values[4] = datetime(2008, 5, 9)
            page.append(values)
        book.save(path)
        book.close()
        self.config['source'].update(path='../master.xlsx', sheet='LAHORE', cell_range='A1:J2')
        self.run_normalize()
        self.assertEqual(len(self.exported()), 1)
        self.assertEqual(self.exported()[0]['DOB'], '2008-05-09')
        inspected = inspect_source(path)
        self.assertTrue(inspected['sheets']['LAHORE']['student_table'])
        self.assertFalse(inspected['sheets']['Overview']['student_table'])

    def test_enrichment_joins_leading_zero_bform_and_conflicts_are_reported(self):
        extra = self.root / 'details.csv'
        self.write_csv(extra, ['BForm', 'Phone'], [{'BForm': '0000001234561', 'Phone': '03001234567'}])
        self.config['enrichment']['sources'] = [{'path': '../details.csv', 'key': 'BForm'}]
        self.run_normalize()
        self.assertEqual(self.exported()[0]['Phone'], '03001234567')
        self.write_csv(extra, ['BForm', 'Phone'], [{'BForm': '0000001234561', 'Phone': '03001234567'}, {'BForm': '0000001234561', 'Phone': '03007654321'}])
        self.run_normalize()
        self.assertIn('enrichment_conflict:enrichment:0', self.provenance()[0]['issues'])

    def test_identity_override_requires_explicit_correction_flag(self):
        path = self.root / 'override.csv'
        headers = ['Roll No', 'DOB', 'IdentityCorrection']
        self.write_csv(path, headers, [{'Roll No': '001234', 'DOB': '01/01/2009', 'IdentityCorrection': ''}])
        self.config['enrichment']['overrides'] = '../override.csv'
        self.run_normalize()
        self.assertEqual(self.exported()[0]['DOB'], '2008-05-09')
        self.assertIn('identity_conflict:DOB', self.provenance()[0]['issues'])
        self.write_csv(path, headers, [{'Roll No': '001234', 'DOB': '01/01/2009', 'IdentityCorrection': 'yes'}])
        self.run_normalize()
        self.assertEqual(self.exported()[0]['DOB'], '2009-01-01')
        self.assertEqual(self.provenance()[0]['origins']['DOB'], 'operator_correction')

    def test_conflicting_scoped_rules_do_not_choose_arbitrarily(self):
        self.config['rules'] = [{'when': {'institutions': ['LAHORE']}, 'set': {'Sport': 'Athletics'}},
                                {'when': {'rolls': ['001234']}, 'set': {'Sport': 'Other'}}]
        self.run_normalize()
        self.assertIn('conflicting_profile_rules:Sport', self.provenance()[0]['issues'])
        self.assertEqual(self.exported()[1]['Sport'], 'Athletics')

    def test_typo_unsupported_mode_and_size_overlap_fail_before_output(self):
        broken = deepcopy(self.config)
        broken['selection']['institutess'] = ['LAHORE']
        self.profile.write_text(json.dumps(broken))
        with self.assertRaisesRegex(ValueError, 'Unknown setting'):
            execute(self.profile)
        self.assertFalse(self.output.exists())
        broken = deepcopy(self.config)
        broken['size_estimation']['rules'][1]['min_age'] = 11
        with self.assertRaisesRegex(ValueError, 'gaps or overlaps'):
            validate(broken)
        broken = deepcopy(self.config)
        broken['gender']['student_cnic_parity']['source_column'] = 'Father CNIC'
        with self.assertRaisesRegex(ValueError, 'student BForm'):
            validate(broken)

    def test_empty_selection_and_publication_failure_preserve_previous_output(self):
        self.run_normalize()
        before = hashes(self.output)
        self.config['selection']['institutions'] = ['No matches']
        with self.assertRaisesRegex(ValueError, 'no saved student'):
            self.run_normalize()
        self.assertEqual(hashes(self.output), before)
        self.config['selection']['institutions'] = ['LAHORE']
        self.config['defaults']['Sport'] = 'Other'
        with patch('bise_normalize.publish.move_staging', side_effect=OSError('simulated rename failure')):
            with self.assertRaises(OSError):
                self.run_normalize()
        self.assertEqual(hashes(self.output), before)
        self.assertFalse(list(self.output.parent.glob('.building-normalized-*')))

    def test_broken_artifact_is_repaired_and_unmanaged_folder_is_preserved(self):
        self.run_normalize()
        (self.output / 'Normalized.csv').write_text('damaged')
        self.assertEqual(execute(self.profile)['status'], 'published')
        self.assertEqual(len(self.exported()), 2)
        (self.output / 'Data/manifest.json').unlink()
        before = hashes(self.output)
        with self.assertRaisesRegex(ValueError, 'unmanaged'):
            execute(self.profile)
        self.assertEqual(hashes(self.output), before)

    def test_source_manifest_mismatch_is_rejected_without_publishing(self):
        (self.root / 'Data').mkdir()
        (self.root / 'Data/manifest.json').write_text(json.dumps({'files': {'master.csv': 'bad hash'}}))
        with self.assertRaisesRegex(ValueError, 'manifest hash'):
            self.run_normalize()
        self.assertFalse(self.output.exists())

    def test_capture_retries_when_source_changes_during_copy(self):
        from bise_normalize import sources
        original = sources.shutil.copyfileobj
        calls = []
        destination = self.root / 'snapshots'
        destination.mkdir()

        def changing(src, dst):
            original(src, dst)
            calls.append(1)
            if len(calls) == 1:
                self.rows.append(student('001239'))
                self.write_source()

        with patch('bise_normalize.sources.shutil.copyfileobj', side_effect=changing):
            result, sha = capture(self.source, destination, 'source')
        self.assertEqual(len(calls), 2)
        self.assertEqual(digest(result), digest(self.source))
        self.assertEqual(sha, digest(self.source))

    def test_profile_init_and_output_lock_are_idempotent_and_exclusive(self):
        before = self.profile.read_bytes()
        self.assertFalse(initialize(self.profile, 'Changed', self.source))
        self.assertEqual(self.profile.read_bytes(), before)
        with locked(self.output):
            with self.assertRaisesRegex(ValueError, 'Another normalizer'):
                execute(self.profile)

    def test_formula_like_names_are_text_and_originals_remain_in_provenance(self):
        self.rows = [student(**{'Student Name': '=1+1'})]
        self.write_source()
        self.run_normalize()
        self.assertEqual(self.exported()[0]['FullName'], "'=1+1")
        self.assertEqual(self.provenance()[0]['data']['FullName'], '=1+1')
        book = load_workbook(self.output / 'Normalized.xlsx')
        self.addCleanup(book.close)
        cell = book['Normalized'].cell(2, COLUMNS.index('FullName') + 1)
        self.assertEqual(cell.data_type, 's')
        self.assertEqual(cell.value, '=1+1')

    def test_portal_name_separators_are_fixed_with_raw_board_value_retained(self):
        self.rows=[student(**{'Father Name':'AHSAN -UL-HAQ'})]
        self.write_source();before=digest(self.source)
        self.run_normalize()
        self.assertEqual(self.exported()[0]['FatherName'],'AHSAN-UL-HAQ')
        info=self.provenance()[0]
        self.assertEqual(info['origins']['FatherName'],'portal_name_separator_formatting')
        self.assertIn('name_formatting:FatherName',info['warnings'])
        self.assertEqual(info['raw']['Father Name'],'AHSAN -UL-HAQ')
        self.assertEqual(digest(self.source),before)

    def test_suggestion_only_mode_keeps_gender_blank_and_no_institute_is_not_an_issue(self):
        self.config['gender']['student_cnic_parity']['mode'] = 'suggestion_only'
        self.run_normalize()
        row = self.provenance()[0]
        self.assertEqual(row['data']['Gender'], '')
        self.assertEqual(row['suggestions']['Gender'], 'Male')
        self.assertNotIn('missing:Institute', row['issues'])

    def test_small_format_helpers_handle_leap_dates_and_invalid_phones(self):
        self.assertEqual(born_on('29/02/2008', '%d/%m/%Y').isoformat(), '2008-02-29')
        with self.assertRaises(ValueError):
            born_on('29/02/2009', '%d/%m/%Y')
        self.assertEqual(clean_phone('923001234567'), '03001234567')
        self.assertEqual(clean_phone('123'), '')
        self.assertEqual(identifier('00000-0123456-2'), '0000001234562')


if __name__ == '__main__':
    unittest.main()
