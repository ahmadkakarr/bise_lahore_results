"""Existing verified cards must gain new fields without another board lookup."""

import asyncio
from argparse import Namespace
from contextlib import redirect_stdout
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from bise_results.presentation import LABELS
FIELD_NAMES = {label: field for field, label in LABELS.items()}

from bise_results.parser import IDENTITY_FIELDS, PARSER_VERSION, parse_result
from bise_results.reports import export_batch, master_name
from bise_results.store import check_receipt, import_batch, jobs_for, locked_store, now, prepare_jobs, reparse_cached_results
from test_identity_fields import FIXTURE, IDENTITY, receipt_for


class CacheUpgradeTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "rolls.csv"
        self.source.write_text("Roll Numbers\n001234\n")
        self.database = self.root / "results.sqlite3"
        self.legacy = parse_result(FIXTURE.read_text(), "001234")
        for key in IDENTITY_FIELDS:
            self.legacy.pop(key)
        self.legacy.update({"parser_version": "ssc-part2-v1", "fetched_at": "2026-10-20T00:00:00+00:00", "source_url": "http://result.biselahore.com/ResultPart2.aspx"})
        self.receipt = receipt_for(self.root, result=self.legacy)
        with locked_store(self.database) as db:
            self.batch = import_batch(db, self.source, "Institution-wise")
            self.job = jobs_for(db, self.batch)[0]["id"]
            with db:
                db.execute("UPDATE jobs SET state='verified',attempts=2,result_json=?,receipt_json=? WHERE id=?", (json.dumps(self.legacy), json.dumps(self.receipt), self.job))

    def test_upgrade_is_idempotent_and_preserves_original_evidence_and_marks(self):
        originals = {name: Path(record["path"]).read_bytes() for name, record in self.receipt.items()}
        with locked_store(self.database) as db:
            report = reparse_cached_results(db, self.batch)
            self.assertEqual(report["upgraded"], 1)
            row = jobs_for(db, self.batch)[0]
            parsed = json.loads(row["result_json"])
            receipt = json.loads(row["receipt_json"])
            self.assertEqual(row["state"], "verified")
            self.assertEqual(row["attempts"], 2)
            self.assertTrue(check_receipt(receipt))
            self.assertEqual(parsed["parser_version"], PARSER_VERSION)
            for key, value in IDENTITY.items():
                self.assertEqual(parsed[key], value)
            for key, value in self.legacy.items():
                if key != "parser_version":
                    self.assertEqual(parsed[key], value, key)
            self.assertEqual(json.loads(Path(receipt["parsed.json"]["path"]).read_text()), parsed)
            self.assertEqual(receipt["card.html"], self.receipt["card.html"])
            self.assertEqual(receipt["card.png"], self.receipt["card.png"])
            self.assertNotEqual(receipt["parsed.json"]["path"], self.receipt["parsed.json"]["path"])
            first_result_json = row["result_json"]
            self.assertEqual(reparse_cached_results(db, self.batch)["already_current"], 1)
            self.assertEqual(jobs_for(db, self.batch)[0]["result_json"], first_result_json)
            self.assertEqual(prepare_jobs(db, self.batch), [])
        for name, before in originals.items():
            self.assertEqual(Path(self.receipt[name]["path"]).read_bytes(), before)

    def test_export_upgrades_both_master_and_institution_files(self):
        with locked_store(self.database) as db:
            destination, summary = export_batch(db, self.batch, self.root / "output")
        self.assertEqual(summary, {"verified": 1, "total": 1})
        groups = json.loads((destination / "Data/institutions.json").read_text())
        for folder, filename in ((destination, f"{master_name(2026)}.xlsx"), (destination / groups[0]["folder"], groups[0]["workbook"])):
            book = load_workbook(folder / filename)
            headers = [FIELD_NAMES[cell.value] for cell in book["All Students"][1]]
            row = dict(zip(headers, next(book["All Students"].iter_rows(min_row=2, values_only=True))))
            for key, value in IDENTITY.items():
                self.assertEqual(row[key], value)
            record = json.loads((folder / "Data/results.json").read_text())[0]
            self.assertNotIn("_receipt", record)
            self.assertEqual(record["obtained"], 551)
        manifest = json.loads((destination / "Data/manifest.json").read_text())
        self.assertTrue(manifest["complete_collection"])
        for name, digest in manifest["files"].items():
            self.assertEqual(hashlib.sha256((destination / name).read_bytes()).hexdigest(), digest)

    def test_collect_uses_cached_card_without_launching_a_browser(self):
        from bise_results.runner import collect
        args = Namespace(csv=self.source, school="Institution-wise", year=2026, database=self.database, retry_unresolved=False, limit=0, workers=1, interval=3, captcha_mode="auto", artifacts=self.root / "artifacts", output=self.root / "output", browser_executable=None)
        with patch("bise_results.runner.async_playwright", side_effect=AssertionError("Browser must not open for a valid cached card")), redirect_stdout(io.StringIO()):
            self.assertEqual(asyncio.run(collect(args)), 0)
        with locked_store(self.database) as db:
            self.assertEqual(json.loads(jobs_for(db, self.batch)[0]["result_json"])["student_cnic"], IDENTITY["student_cnic"])
            self.assertEqual(jobs_for(db, self.batch)[0]["attempts"], 2)

    def test_corrupt_evidence_is_not_used_for_backfill(self):
        (self.root / "card.html").write_text("tampered")
        with locked_store(self.database) as db:
            self.assertEqual(reparse_cached_results(db, self.batch)["unavailable"], 1)
            row = jobs_for(db, self.batch)[0]
            self.assertEqual(json.loads(row["result_json"]), self.legacy)
            self.assertEqual(len(prepare_jobs(db, self.batch)), 1)

    def test_parser_failure_marks_review_instead_of_trusting_old_cache(self):
        source = FIXTURE.read_text().replace("2026", "2025")
        receipt = receipt_for(self.root, source=source, result=self.legacy)
        with locked_store(self.database) as db:
            with db:
                db.execute("UPDATE jobs SET receipt_json=? WHERE id=?", (json.dumps(receipt), self.job))
            self.assertEqual(reparse_cached_results(db, self.batch)["needs_review"], 1)
            self.assertEqual(jobs_for(db, self.batch)[0]["state"], "needs_review")

    def test_failed_artifact_write_preserves_original_cache(self):
        with locked_store(self.database) as db:
            with patch("bise_results.store._write_reparsed_artifacts", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    reparse_cached_results(db, self.batch)
            row = jobs_for(db, self.batch)[0]
            self.assertEqual(json.loads(row["result_json"]), self.legacy)
            self.assertEqual(json.loads(row["receipt_json"]), self.receipt)
            self.assertTrue(check_receipt(self.receipt))

    def test_incomplete_current_version_is_also_upgraded(self):
        incomplete = {**self.legacy, "parser_version": PARSER_VERSION}
        receipt = receipt_for(self.root, result=incomplete)
        with locked_store(self.database) as db:
            with db:
                db.execute("UPDATE jobs SET result_json=?,receipt_json=? WHERE id=?", (json.dumps(incomplete), json.dumps(receipt), self.job))
            self.assertEqual(reparse_cached_results(db, self.batch)["upgraded"], 1)
            parsed = json.loads(jobs_for(db, self.batch)[0]["result_json"])
            self.assertEqual(parsed["father_cnic"], IDENTITY["father_cnic"])

    def review_attempt(self, db, source):
        folder = self.root / "rejected-response"
        folder.mkdir()
        (folder / "card.html").write_text(source)
        from PIL import Image
        Image.new("RGB", (2, 2), "white").save(folder / "card.png")
        with db:
            db.execute("UPDATE jobs SET state='needs_review',error='Old parser rejected layout',result_json=NULL,receipt_json=NULL WHERE id=?", (self.job,))
            db.execute("INSERT INTO attempts (job_id,started_at,ended_at,outcome,evidence) VALUES (?,?,?,?,?)", (self.job, now(), now(), "needs_review", str(folder)))
        return folder

    def test_previously_reviewed_cards_recover_without_another_attempt(self):
        for name in ("absent", "not-improved", "pass-prefix"):
            with self.subTest(layout=name), tempfile.TemporaryDirectory() as separate:
                root = Path(separate)
                source = (FIXTURE.parent / f"ssc-part2-{name}.html").read_text()
                # A separate store for each real layout prevents cached results
                # from hiding a later fixture's recovery path.
                with locked_store(root / "results.sqlite3") as db:
                    batch = import_batch(db, self.source, "Test")
                    old_root = self.root
                    self.root = root
                    try:
                        folder = self.review_attempt(db, source)
                    finally:
                        self.root = old_root
                    original = (folder / "card.html").read_bytes()
                    count = db.execute("SELECT count(*) FROM attempts").fetchone()[0]
                    report = reparse_cached_results(db, batch)
                    self.assertEqual(report["recovered"], 1)
                    row = jobs_for(db, batch)[0]
                    self.assertEqual(row["state"], "verified")
                    self.assertIsNone(row["error"])
                    self.assertTrue(check_receipt(json.loads(row["receipt_json"])))
                    self.assertEqual(json.loads(row["result_json"])["name"], "Sample Candidate")
                    self.assertEqual((folder / "card.html").read_bytes(), original)
                    self.assertEqual(db.execute("SELECT count(*) FROM attempts").fetchone()[0], count)
                    self.assertEqual(prepare_jobs(db, batch, retry_unresolved=True), [])
                    self.assertEqual(reparse_cached_results(db, batch)["already_current"], 1)
                    destination, summary = export_batch(db, batch, root / "output")
                    self.assertEqual(summary, {"verified": 1, "total": 1})
                    book = load_workbook(destination / f"{master_name(2026)}.xlsx")
                    self.assertEqual(book["Missing Results"].max_row, 1)
                    if name != "pass-prefix":
                        headers = [FIELD_NAMES[cell.value] for cell in book["All Students"][1]]
                        self.assertIsNone(book["All Students"].cell(2, headers.index("obtained") + 1).value)

    def test_reviewed_captcha_page_cannot_become_a_student_record(self):
        with locked_store(self.database) as db:
            self.review_attempt(db, "<html><body>Invalid Captcha. Please try again.</body></html>")
            self.assertEqual(reparse_cached_results(db, self.batch)["recovered"], 0)
            self.assertEqual(jobs_for(db, self.batch)[0]["state"], "needs_review")

    def test_unavailable_review_capture_is_left_unresolved(self):
        with locked_store(self.database) as db:
            with db:
                db.execute("UPDATE jobs SET state='needs_review',result_json=NULL,receipt_json=NULL WHERE id=?", (self.job,))
            self.assertEqual(reparse_cached_results(db, self.batch)["unavailable"], 1)
            self.assertEqual(jobs_for(db, self.batch)[0]["state"], "needs_review")


if __name__ == "__main__":
    unittest.main()
