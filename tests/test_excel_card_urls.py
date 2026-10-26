"""Card URL exclusion beyond Excel's per-sheet hyperlink limit."""
from pathlib import Path
import tempfile
import unittest
import warnings
import zipfile
from openpyxl import load_workbook
from bise_results.fast_excel import build_workbook
from bise_results.presentation import build_workbook as legacy_workbook
from bise_results.reports import safe_text


class ExcelCardUrlTests(unittest.TestCase):
    def test_more_than_65530_students_have_no_card_urls_or_url_limit_warnings(self):
        count=65531
        records=[{'roll':f'{i:06d}','name':'Synthetic Student','state':'verified','institution':'Synthetic School'} for i in range(count)]
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'large.xlsx'
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter('always')
                build_workbook(path,records,{'verified':count,'total':count},'Large regression',
                    ['roll','name','result_card'],safe_text,{'roll'})
            self.assertFalse(any('URL' in str(w.message) for w in caught))
            with zipfile.ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
                for name in archive.namelist():
                    if name.startswith('xl/') and (name.endswith('.rels') or name=='xl/sharedStrings.xml'):
                        self.assertNotIn(b'.png',archive.read(name),name)
            book=load_workbook(path,read_only=True)
            try:
                for name in ('All Students','Synthetic School'):
                    page=book[name]
                    self.assertEqual(page.max_row,count+1)
                    self.assertEqual([cell.value for cell in next(page.rows)],['Roll No','Student Name'])
                self.assertEqual(book['All Students'].cell(2,1).value,'000000')
            finally:book.close()

    def test_legacy_writer_also_omits_card_column_and_links(self):
        records=[{'roll':'001234','name':'Synthetic','state':'verified','institution':'Synthetic School'}]
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'legacy.xlsx'
            legacy_workbook(path,records,{'verified':1,'total':1},'Legacy',
                ['roll','name','result_card'],safe_text,{'roll'})
            book=load_workbook(path)
            try:
                for name in ('All Students','Synthetic School'):
                    self.assertEqual([cell.value for cell in book[name][1]],['Roll No','Student Name'])
                    self.assertTrue(all(cell.hyperlink is None for row in book[name] for cell in row))
                self.assertIsNotNone(book['Overview']['A5'].hyperlink)
            finally:book.close()
