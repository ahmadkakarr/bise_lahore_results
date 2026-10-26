"""Correctness checks against a de-identified table captured from the live board."""

import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from PIL import Image

from bise_results.parser import parse_result, ReviewRequired
from bise_results.reports import export_batch, master_name
from bise_results.store import check_receipt, import_batch, jobs_for, locked_store, now, prepare_jobs

FIXTURE = Path(__file__).parent / "fixtures/ssc-part2.html"


class ParserTests(unittest.TestCase):
    def test_live_table_reconciles_and_preserves_improvement(self):
        result = parse_result(FIXTURE.read_text(), "001234")
        self.assertEqual((result["obtained"], result["maximum"], len(result["subjects"])), (551, 1100, 8))
        self.assertIn("MARKS IMPROVED", result["total_raw"])
        self.assertEqual(result["roll"], "001234")
        self.assertEqual(result["subjects"][0]["practical_ii_raw"], "--")

    def test_identity_mismatch_needs_review(self):
        original = FIXTURE.read_text()
        for source, roll in [(original, "999999"), (original.replace("2026", "2025"), "001234"), (original.replace("Sample Candidate", ""), "001234")]:
            with self.subTest(roll=roll), self.assertRaises(ReviewRequired):
                parse_result(source, roll)

    def test_withheld_marks_stay_null(self):
        source = FIXTURE.read_text().replace("MARKS OBTAINED:551 C MARKS IMPROVED", "RESULT WITHHELD")
        source = source.replace("The candidate has Passed and obtained Marks Five Hundred and Fifty One.", "Result withheld.")
        result = parse_result(source, "001234")
        self.assertIsNone(result["obtained"])
        self.assertIsNone(result["percentage"])
        self.assertEqual(result["total_raw"], "RESULT WITHHELD")


class StoreAndExportTests(unittest.TestCase):
    def test_resume_membership_and_institution_reports(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "rolls.csv"
            source.write_text("Roll Numbers\n001234\n001234\n567890\n")
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            with locked_store(root / "results.sqlite3") as db:
                batch = import_batch(db, source, "Input A")
                self.assertEqual(batch, import_batch(db, source, "Input A"))
                other = import_batch(db, source, "Input B")
                self.assertNotEqual(batch, other)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0], 2)
                self.assertEqual(len(jobs_for(db, batch)), 2)
                first, second = jobs_for(db, batch)
                # A crashed process left one running claim and unfinished attempt.
                with db:
                    db.execute("UPDATE jobs SET state='running' WHERE id=?", (first["id"],))
                    db.execute("INSERT INTO attempts (job_id,started_at) VALUES (?,?)", (first["id"], now()))
                self.assertEqual(len(prepare_jobs(db, batch)), 2)
                self.assertEqual(db.execute("SELECT outcome FROM attempts").fetchone()[0], "interrupted")
                parsed = parse_result(FIXTURE.read_text(), "001234")
                parsed["name"] = "=1+1"  # Board text must not execute as an Excel formula.
                (root / "card.html").write_text(FIXTURE.read_text())
                Image.new("RGB", (2, 2), "white").save(root / "card.png")
                (root / "parsed.json").write_text(json.dumps(parsed))
                receipt = {name: {"path": str(root / name), "sha256": hashlib.sha256((root / name).read_bytes()).hexdigest()} for name in ("card.html", "card.png", "parsed.json")}
                with db:
                    db.execute("UPDATE jobs SET state='verified',result_json=?,receipt_json=? WHERE id=?", (json.dumps(parsed), json.dumps(receipt), first["id"]))
                self.assertEqual([job["id"] for job in prepare_jobs(db, batch)], [second["id"]])
                self.assertTrue(check_receipt(receipt))
                destination, summary = export_batch(db, batch, root / "output")
                self.assertEqual(summary, {"verified": 1, "pending": 1, "total": 2})
                groups = json.loads((destination / "Data/institutions.json").read_text())
                self.assertEqual(groups[0]["institution_district"], parsed["institution"])
                book = load_workbook(destination / f"{master_name(2026)}.xlsx", data_only=False)
                self.assertEqual(book["All Students"].max_row, 3)  # deduplicated membership
                self.assertEqual(book["All Students"]["A2"].value, "001234")
                self.assertEqual(book["All Students"]["C2"].value, "'=1+1")
                self.assertEqual(book["All Students"]["C2"].data_type, "s")
                manifest = json.loads((destination / "Data/manifest.json").read_text())
                self.assertFalse(manifest["complete_collection"])
                for name, digest in manifest["files"].items():
                    self.assertEqual(hashlib.sha256((destination / name).read_bytes()).hexdigest(), digest)
                # Corrupted evidence must invalidate the verified cache and requeue.
                (root / "card.png").write_bytes(b"corrupted")
                self.assertFalse(check_receipt(receipt))
                self.assertEqual(len(prepare_jobs(db, batch)), 2)
                import_batch(db, source, "Input A", year=2025)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0], 4)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)

    def test_process_lock_blocks_a_second_runner(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.sqlite3"
            with locked_store(path):
                with self.assertRaises(RuntimeError):
                    with locked_store(path):
                        self.fail("Second runner acquired the lock")


if __name__ == "__main__":
    unittest.main()
