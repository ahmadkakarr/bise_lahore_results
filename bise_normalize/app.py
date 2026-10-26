"""Snapshot-to-output orchestration; no collector or browser dependency."""

from collections import Counter
from contextlib import nullcontext
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile

from . import COLUMNS, VERSION
from .config import load, resolve
from .engine import prepare
from .publish import build, current, locked
from .sources import snapshots, table, digest
from . import state
from .portal import evaluate
from youthgames_shared.catalog import load as load_catalog, signature as content_signature


def inputs(config, path):
    result = {'source': resolve(path, config['source']['path'])}
    if config['target']['template']:
        result['template'] = resolve(path, config['target']['template'])
    for i, entry in enumerate(config['enrichment']['sources']):
        result[f'enrichment:{i}'] = resolve(path, entry['path'])
    if config['enrichment']['overrides']:
        result['overrides'] = resolve(path, config['enrichment']['overrides'])
    if config['portal']['enabled']:
        result['catalog'] = resolve(path, config['validation']['catalog'])
    return result


def fingerprint(config, input_hashes):
    effective = deepcopy(config)
    effective['source'].pop('path')
    effective['output'].pop('folder')
    effective['target'].pop('template')
    for entry in effective['enrichment']['sources']:
        entry.pop('path')
    effective['enrichment']['overrides'] = bool(effective['enrichment']['overrides'])
    return hashlib.sha256(json.dumps({'version': VERSION, 'config': effective, 'inputs': input_hashes}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def execute(path, mode='run', preview_limit=20):
    path = Path(path).resolve()
    config = load(path)
    files = inputs(config, path)
    destination = resolve(path, config['output']['folder'])
    portal = config['portal']
    ledger_path = resolve(path, portal['normalization_database'])
    registration_path = resolve(path, portal['registration_database'])
    if portal['enabled'] and (ledger_path == registration_path or ledger_path.is_relative_to(destination) or registration_path.is_relative_to(destination)):
        raise ValueError('State databases must be separate and outside generated output')
    for source in [path, *files.values()]:
        if source == destination or source.is_relative_to(destination):
            raise ValueError('Output folder must not contain source/config/enrichment files')
    if config['source']['format'] == 'auto':
        config['source']['format'] = files['source'].suffix.lower().removeprefix('.')
        if config['source']['format'] in ('sqlite3', 'db'):
            config['source']['format'] = 'sqlite'
    context = locked(destination) if mode == 'run' else nullcontext()
    with context, snapshots(files, source_settings=config['source']) as captured:
        if 'template' in captured:
            _, headers = next(table(captured['template'][0], {'format': 'csv', 'header_row': 1, 'sheet': None, 'cell_range': None}))
            if headers != COLUMNS:
                raise ValueError('Sample headers differ from the exact target column contract')
        input_hashes = {k: v[1] for k, v in captured.items()}
        catalog, states = None, {}
        if portal['enabled']:
            catalog = load_catalog(captured['catalog'][0], portal['max_catalog_age_hours'])
            states = state.registration_snapshot(registration_path, portal['campaign'])
            if config['source'].get('allocation_scheme'):
                states=state.allocated_states(states,files['source'],config['source']['allocation_scheme'])
            # Refresh timestamps do not change the meaning of a catalog or registration state.
            input_hashes['catalog'] = catalog['content_sha256']
            input_hashes['registration_states'] = content_signature(states)
        signature = fingerprint(config, input_hashes)
        ledger_context = state.opened(ledger_path, state.transform_context(config, captured)) if portal['enabled'] and mode == 'run' else nullcontext(None)
        with ledger_context as ledger, tempfile.TemporaryDirectory(prefix='bise-normalize-work-') as work:
            if mode == 'run':
                previous = current(destination, signature, config['profile'])
                indexed = not ledger or ledger.db.execute('SELECT 1 FROM runs WHERE profile=? AND fingerprint=? AND manifest_hash=?',
                    (config['profile'], signature, digest(destination / 'Data/manifest.json') if previous else '')).fetchone()
                if previous and indexed:
                    return {'status': 'already_up_to_date', 'output': str(destination), 'fingerprint': signature, 'summary': previous['summary'],
                            'cache': {'reused': 0, 'transformed': 0}}
            prepared = prepare(captured['source'][0], config, captured, Path(work) / 'work.sqlite3', cache=ledger)
            try:
                if catalog:
                    evaluate(prepared, catalog, states)
                result = {'status': mode, 'output': str(destination), 'fingerprint': signature,
                          'summary': dict(prepared.summary), 'source_sha256': input_hashes['source']}
                if mode == 'run':
                    if not previous:
                        build(destination, prepared, config, signature, input_hashes)
                    if ledger:
                        ledger.published(prepared, config, signature, digest(destination / 'Data/manifest.json'))
                        result['cache'] = {'reused': ledger.hits, 'transformed': ledger.transformed}
                    result['status'] = 'published'
                elif mode == 'preview':
                    from itertools import islice
                    result['preview'] = [{'roll': r['roll'], 'source_rows': r['source_rows'], 'issues': r['issues'], 'warnings': r['warnings']} for r in islice(prepared.records(), preview_limit)]
                return result
            finally:
                prepared.close()


def inspect_source(source, sample=None):
    source = Path(source).resolve()
    if source.suffix.lower() in ('.sqlite3', '.sqlite', '.db'):
        import sqlite3
        db = sqlite3.connect(source.as_uri() + '?mode=ro', uri=True)
        try:
            return {'source': str(source), 'format': 'sqlite', 'exam_scopes': [dict(zip(('year','course','exam_type','state','records'), row))
                for row in db.execute('SELECT year,course,exam_type,state,count(*) FROM jobs GROUP BY year,course,exam_type,state')]}
        finally:
            db.close()
    paths = {'source': source}
    if sample:
        paths['template'] = Path(sample).resolve()
    result = {'source': str(source)}
    with snapshots(paths) as captured:
        if source.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(captured['source'][0], read_only=True, data_only=True)
            try:
                result['sheets'] = {page.title: {'rows': page.max_row, 'columns': page.max_column,
                    'student_table': {'Roll No', 'Student Name'}.issubset(set(next(page.iter_rows(max_row=1, values_only=True), ())))} for page in book}
            finally:
                book.close()
            return result
        reader = iter(table(captured['source'][0], {'format': 'csv', 'sheet': None, 'cell_range': None, 'header_row': 1}))
        _, result['headers'] = next(reader)
        labels, states, total = Counter(), Counter(), 0
        for _, row in reader:
            total += 1
            labels[row.get('Institution / District', '')] += 1
            states[row.get('Collection Status', '')] += 1
        result.update({'rows': total, 'institutions': dict(labels), 'collection_states': dict(states)})
        if 'template' in captured:
            _, headers = next(table(captured['template'][0], {'format': 'csv', 'sheet': None, 'cell_range': None, 'header_row': 1}))
            result['target_headers'] = headers
            result['target_schema_matches'] = headers == COLUMNS
    return result
