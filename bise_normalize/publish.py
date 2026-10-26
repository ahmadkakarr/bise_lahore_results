"""Friendly artifacts, verified reruns, profile locks and rollback publication."""

from contextlib import contextmanager
from datetime import datetime, timezone
import csv
import fcntl
import hashlib
import json
from pathlib import Path
import shutil
import tempfile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from . import COLUMNS, VERSION
from .sources import digest


def hashes(folder):
    return {p.relative_to(folder).as_posix(): digest(p) for p in sorted(folder.rglob('*')) if p.is_file()}


def state_folder(destination):
    return destination.parent / '.normalization'


@contextmanager
def locked(destination):
    root = state_folder(destination) / 'locks'
    root.mkdir(parents=True, exist_ok=True)
    name = hashlib.sha256(str(destination).encode()).hexdigest() + '.lock'
    with (root / name).open('a') as stream:
        try:
            fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ValueError('Another normalizer owns this output folder') from exc
        try:
            yield
        finally:
            fcntl.flock(stream, fcntl.LOCK_UN)


def current(destination, fingerprint, profile):
    manifest_path = destination / 'Data/manifest.json'
    if not destination.exists():
        return None
    if not manifest_path.is_file():
        raise ValueError('Output contains unmanaged files; choose a separate normalization folder')
    manifest = json.loads(manifest_path.read_text())
    if manifest.get('application') != 'bise_normalize' or manifest.get('profile') != profile:
        raise ValueError('Output belongs to another application/profile; choose another folder')
    if manifest.get('fingerprint') == fingerprint:
        actual = hashes(destination)
        actual.pop('Data/manifest.json', None)
        if actual == manifest.get('files'):
            return manifest
    return None


def csv_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def csv_file(path, columns, values):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(columns)
        writer.writerows([csv_safe(v) for v in row] for row in values)


def sheet(book, title, headers):
    page = book.create_sheet(title)
    page.freeze_panes = 'D2'
    for index, heading in enumerate(headers, 1):
        from openpyxl.utils import get_column_letter
        page.column_dimensions[get_column_letter(index)].width = 24 if heading not in ('Issues', 'Origins', 'Warnings', 'Address', 'Institute') else 55
    row = []
    for heading in headers:
        cell = WriteOnlyCell(page, value=heading)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1C4D24')
        row.append(cell)
    page.append(row)
    return page


def excel_row(page, values):
    cells = []
    for value in values:
        if isinstance(value, str):
            cell = WriteOnlyCell(page, value=ILLEGAL_CHARACTERS_RE.sub('', value))
            cell.data_type = 's'
            cell.number_format = '@'
        else:
            cell = WriteOnlyCell(page, value=value)
        cells.append(cell)
    page.append(cells)


def workbook(path, prepared, profile, missing_only=False):
    book = Workbook(write_only=True)
    if missing_only:
        page = sheet(book, 'Missing Details', ['Roll No', 'Source Rows', 'Issues', *COLUMNS])
        for row in prepared.records(only_missing=True):
            excel_row(page, [row['roll'], ', '.join(map(str, row['source_rows'])), '; '.join(row['issues']), *[row['data'][k] for k in COLUMNS]])
    else:
        overview = sheet(book, 'Overview', ['Metric', 'Value'])
        excel_row(overview, ['Profile', profile])
        for key, value in sorted(prepared.summary.items()):
            excel_row(overview, [key, value])
        excel_row(overview, ['Meaning of Complete', 'Required profile fields present; includes flagged derived/estimated defaults.'])
        excel_row(overview, ['Portal eligibility', 'Checked against captured official catalog' if hasattr(prepared, 'catalog') else 'Not checked'])
        page = sheet(book, 'Normalized', COLUMNS)
        details = sheet(book, 'Details', ['Roll No', 'Source Rows', 'Board Institution / District', 'Age', 'Issues', 'Warnings', 'Origins'])
        for row in prepared.records():
            excel_row(page, [row['data'][k] for k in COLUMNS])
            excel_row(details, [row['roll'], ', '.join(map(str, row['source_rows'])), row['raw'].get('Institution / District', ''), row['age'], '; '.join(row['issues']), '; '.join(row['warnings']), json.dumps(row['origins'], ensure_ascii=False, sort_keys=True)])
        if hasattr(prepared, 'catalog'):
            from .portal import SUBSETS, subset
            for name in SUBSETS:
                page = sheet(book, name, ['Roll No', 'Age Reason', 'Registration State', 'Portal Issues', *COLUMNS])
                for row in subset(prepared, name):
                    excel_row(page, [row['roll'], row['portal']['age']['reason'], row['portal']['registration_state'], '; '.join(row['portal']['issues']), *[row['data'][k] for k in COLUMNS]])
    book.save(path)


def move_staging(staging, destination):
    staging.rename(destination)


def build(destination, prepared, config, fingerprint, input_hashes):
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix='.building-normalized-', dir=destination.parent))
    previous = None
    try:
        (staging / 'Data').mkdir()
        csv_file(staging / 'Normalized.csv', COLUMNS, ([row['data'][k] for k in COLUMNS] for row in prepared.records()))
        if config['output']['write_complete_subset']:
            csv_file(staging / 'Complete Records.csv', COLUMNS, ([row['data'][k] for k in COLUMNS] for row in prepared.records(only_complete=True)))
        if config['output']['write_excel']:
            workbook(staging / 'Normalized.xlsx', prepared, config['profile'])
            workbook(staging / 'Missing Details.xlsx', prepared, config['profile'], missing_only=True)
        csv_file(staging / 'Data/excluded.csv', ['Source Row', 'Roll No', 'Reason'], prepared.db.execute('SELECT * FROM excluded ORDER BY source_row'))
        csv_file(staging / 'Data/duplicates.csv', ['Kind', 'Key', 'Roll No', 'Source Row'], prepared.db.execute('SELECT * FROM duplicates ORDER BY kind,key,roll,source_row'))
        if config['output']['write_provenance']:
            with (staging / 'Data/provenance.jsonl').open('w', encoding='utf-8') as stream:
                for row in prepared.records():
                    stream.write(json.dumps(row, sort_keys=True, ensure_ascii=False) + '\n')
        if hasattr(prepared, 'catalog'):
            from .portal import SUBSETS, subset
            report_headers = ['Roll No', 'Board Institution / District', 'Reason', 'Cutoff Date', 'Minimum DOB', 'Maximum DOB', 'Registration State', 'Details', *COLUMNS]
            for name in SUBSETS:
                exact_schema = name in ('Age Eligible', 'Eligible', 'Ready to Register')
                def report_rows():
                    for row in subset(prepared, name):
                        values = [row['data'][k] for k in COLUMNS]
                        age = row['portal']['age']; rule = age.get('rule', {})
                        yield values if exact_schema else [row['roll'], row['raw'].get('Institution / District', ''), age['reason'], rule.get('cutoff_date', ''), rule.get('min_dob', ''), rule.get('max_dob', ''), row['portal']['registration_state'], '; '.join(row['issues'] + row['portal']['issues']), *values]
                csv_file(staging / (name + '.csv'), COLUMNS if exact_schema else report_headers, report_rows())
                with (staging / 'Data' / (name.lower().replace(' ', '-') + '-provenance.jsonl')).open('w', encoding='utf-8') as stream:
                    for row in subset(prepared, name):
                        stream.write(json.dumps(row, sort_keys=True, ensure_ascii=False) + '\n')
            (staging / 'Data/portal-catalog.json').write_text(json.dumps(prepared.catalog, indent=2, ensure_ascii=False) + '\n')
        (staging / 'Data/effective-profile.json').write_text(json.dumps(config, indent=2, ensure_ascii=False) + '\n')
        summary = dict(prepared.summary)
        (staging / 'Summary.txt').write_text(
            f"Profile: {config['profile']}\nNormalized students: {summary['normalized']}\n"
            f"Complete against profile: {summary['complete']}\nMissing/conflicting details: {summary['missing_details']}\n"
            "Normalized.csv contains all students; Data/provenance.jsonl retains field origins.\n"
            "Normalized.csv uses the exact 17-column sample format, including partial records.\n"
            "Complete means required values are present; derived gender, shared contacts, board-label addresses and estimated sizes remain flagged in Details/provenance.\n" +
            (f"Age eligible: {summary['age_eligible']}\nAge ineligible: {summary['age_ineligible']}\nAge unknown: {summary['age_unknown']}\n"
             f"Eligible with accepted fields: {summary['eligible']}\nReady to register: {summary['ready_to_register']}\nAlready registered: {summary['already_registered']}\n"
             "Use Ready to Register.csv for upload; Ineligible.csv explains age exclusions.\n"
             "Ready excludes known registered/uncertain/in-progress/rejected entries at snapshot time. Uploader rechecks its database before submitting.\n"
             if hasattr(prepared, 'catalog') else "Portal eligibility was not checked.\n") +
            "No live registration or BISE collection was performed.\n"
            "Edit the profile or override file, then rerun. Generated files may be replaced.\n",
            encoding='utf-8')
        manifest = {'application': 'bise_normalize', 'version': VERSION, 'profile': config['profile'],
                    'fingerprint': fingerprint, 'input_hashes': input_hashes, 'summary': summary,
                    'field_origin_counts': dict(prepared.origins), 'live_rules_checked': hasattr(prepared, 'catalog'),
                    'exported_at': datetime.now(timezone.utc).isoformat(), 'files': hashes(staging)}
        (staging / 'Data/manifest.json').write_text(json.dumps(manifest, indent=2, sort_keys=True) + '\n')
        if destination.exists():
            # current() has already checked ownership under the output lock.
            old_hashes = hashes(destination)
            archive = state_folder(destination) / 'previous' / hashlib.sha256(str(destination).encode()).hexdigest()[:16]
            archive.mkdir(parents=True, exist_ok=True)
            previous_root = Path(tempfile.mkdtemp(prefix='version-', dir=archive))
            previous = previous_root / 'Files'
            destination.rename(previous)
            try:
                if hashes(previous) != old_hashes:
                    raise ValueError('Previous-output inventory verification failed')
                (previous_root / 'Inventory.json').write_text(json.dumps(old_hashes, sort_keys=True))
            except BaseException:
                previous.rename(destination)
                previous = None
                raise
        try:
            move_staging(staging, destination)
        except BaseException:
            if previous is not None and not destination.exists():
                previous.rename(destination)
            raise
        return manifest
    finally:
        if staging.exists():
            shutil.rmtree(staging)
