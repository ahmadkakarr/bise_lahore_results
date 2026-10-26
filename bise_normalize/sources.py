"""Coherent read-only snapshots, CSV/XLSX student tables and selections."""

from contextlib import contextmanager
from datetime import date, datetime
import csv
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import sqlite3

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def stamp(path):
    stat = Path(path).stat()
    return stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns


def capture(path, folder, name):
    path = Path(path)
    target = folder / (name + path.suffix.lower())
    for _ in range(3):
        try:
            before = stamp(path)
            with path.open('rb') as src, target.open('wb') as dst:
                shutil.copyfileobj(src, dst)
            sha = digest(target)
            if stamp(path) != before:
                continue
            metadata = path.parent / 'Data/manifest.json'
            if metadata.is_file():
                manifest = json.loads(metadata.read_text())
                expected = manifest.get('files', {}).get(path.name)
                if expected is not None and sha != expected:
                    continue
            if stamp(path) != before:
                continue
            return target, sha
        except FileNotFoundError:
            continue
    raise ValueError(f'Input changed during snapshot or manifest hash did not match: {path.name}; retry when publication finishes')


def sqlite_table(path, settings):
    """Project saved exam identities from a coherent read-only WAL transaction."""
    if settings.get('sheet') is not None or settings.get('cell_range') is not None or settings.get('header_row', 1) != 1:
        raise ValueError('SQLite source has no worksheets/cell ranges; use selection filters')
    exam = settings.get('exam', {})
    if (exam.get('board') != 'BISE Lahore' or exam.get('course') not in ('SSC','HSSC')
            or exam.get('part') not in ('I', 'II') or exam.get('session') != 'Annual'
            or (exam['course']=='HSSC' and exam['part']!='II')):
        raise ValueError('SQLite source requires an explicit BISE Lahore SSC/HSSC Annual exam scope')
    from bise_results.reports import RESULT_COLUMNS
    from bise_results.presentation import LABELS, display
    db = sqlite3.connect(Path(path).resolve().as_uri() + '?mode=ro', uri=True, timeout=10)
    db.row_factory = sqlite3.Row
    try:
        db.execute('BEGIN')
        scope = (exam['year'], exam['course'], '1' if exam['part'] == 'I' else '2')
        allocation = settings.get('allocation_scheme')
        if not allocation and any(row['name'] == 'allocation_scheme' for row in db.execute('PRAGMA table_info(jobs)')):
            raise ValueError('Allocated databases require source.allocation_scheme; cohorts cannot be combined implicitly')
        extra = ' AND allocation_scheme=?' if allocation else ''
        rows = db.execute("SELECT roll,result_json FROM jobs WHERE year=? AND course=? AND exam_type=? AND state='verified'" + extra + ' ORDER BY roll',
                          scope + ((allocation,) if allocation else ()))
        fields = ['Division','District','Tehsil','UC','Sport','Administrative Allocation','Sport Allocation'] if allocation else []
        yield None, [LABELS[key] for key in RESULT_COLUMNS] + fields
        for number, row in enumerate(rows, 2):
            record = json.loads(row['result_json'])
            if record.get('roll') != row['roll']:
                raise ValueError('Stored result identity differs from its queue key')
            if str(record.get('exam_type', '1' if exam['part'] == 'I' else '2')) != ('1' if exam['part'] == 'I' else '2') or record.get('year', exam['year']) != exam['year']:
                raise ValueError('Stored result differs from the requested exam scope')
            record['state'] = 'verified'
            # No generated Result Cards folder is required for DB normalization.
            output = {LABELS[key]: text(display(record, key)) if key != 'result_card' else '' for key in RESULT_COLUMNS}
            if allocation:
                metadata = record.get('administrative_allocation', {})
                if metadata.get('scheme') != allocation or metadata.get('residence_inferred') is not False:
                    raise ValueError('Allocated source lacks the explicit administrative assignment provenance')
                output.update({key: text(record.get(key)) for key in fields if key not in ('Administrative Allocation','Sport Allocation')})
                output['Administrative Allocation'] = json.dumps(metadata, sort_keys=True)
                output['Sport Allocation'] = json.dumps(record['sport_allocation'],sort_keys=True) if record.get('sport_allocation') else ''
            yield number, output
    finally:
        db.close()


def capture_sqlite(path, folder, name, settings):
    # Snapshot only the selected exam's saved data, never copy a live DB file
    # without its WAL. Stable projection hashes ignore unrelated queue changes.
    target = Path(folder) / (name + '.csv')
    rows = iter(sqlite_table(path, settings))
    _, headers = next(rows)
    try:
        with target.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            for _, row in rows:
                writer.writerow(row)
    finally:
        rows.close()
    return target, digest(target)


@contextmanager
def snapshots(paths, source_settings=None):
    with tempfile.TemporaryDirectory(prefix='bise-normalize-input-') as directory:
        folder = Path(directory)
        result = {}
        for i, (key, value) in enumerate(paths.items()):
            database = key == 'source' and source_settings and (source_settings['format'] == 'sqlite' or
                source_settings['format'] == 'auto' and Path(value).suffix.lower() in ('.sqlite3', '.sqlite', '.db'))
            result[key] = capture_sqlite(value, folder, str(i), source_settings) if database else capture(value, folder, str(i))
        yield result


def headers_valid(headers):
    if not headers or any(not h for h in headers) or len(headers) != len(set(headers)):
        raise ValueError('Input headers must be nonblank and unique')


def table(path, settings):
    kind = settings['format']
    if kind == 'auto':
        kind = Path(path).suffix.lower().removeprefix('.')
    if kind in ('sqlite', 'sqlite3', 'db'):
        if Path(path).suffix.lower() != '.csv':
            yield from sqlite_table(path, settings)
            return
        kind = 'csv'  # Coherent exam projection made by capture_sqlite().
    if kind not in ('csv', 'xlsx'):
        raise ValueError('Source must be CSV or XLSX')
    if kind == 'csv':
        if settings['sheet'] is not None or settings['cell_range'] is not None:
            raise ValueError('CSV has no sheets/cell ranges; use selection.row_ranges')
        with Path(path).open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream)
            headers = None
            for number, values in enumerate(reader, 1):
                if number < settings['header_row']:
                    continue
                if headers is None:
                    headers = [text(v) for v in values]
                    headers_valid(headers)
                    yield None, headers
                    continue
                if not any(text(v) for v in values):
                    continue
                if len(values) != len(headers):
                    raise ValueError(f'CSV record {number} has {len(values)} fields; expected {len(headers)}')
                yield number, dict(zip(headers, map(text, values)))
            if headers is None:
                raise ValueError('Source has no header row')
        return
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = settings['sheet'] or 'All Students'
        if sheet_name not in book.sheetnames:
            raise ValueError(f'Worksheet not found: {sheet_name}')
        sheet = book[sheet_name]
        bounds = range_boundaries(settings['cell_range']) if settings['cell_range'] else (1, settings['header_row'] + 1, sheet.max_column, sheet.max_row)
        first_col, first_row, last_col, last_row = bounds
        if any(v is None or v < 1 for v in bounds):
            raise ValueError('Use a bounded A1 cell range, e.g. A1:V500')
        header_row = settings['header_row']
        values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, min_col=first_col, max_col=last_col, values_only=True))
        headers = [text(v) for v in values]
        headers_valid(headers)
        yield None, headers
        begin = max(first_row, header_row + 1)
        for number, values in enumerate(sheet.iter_rows(min_row=begin, max_row=last_row, min_col=first_col, max_col=last_col, values_only=True), begin):
            if any(v is not None for v in values):
                yield number, dict(zip(headers, map(text, values)))
    finally:
        book.close()


def selected(row, number, settings):
    if settings['collection_states'] and row.get('Collection Status', '') not in settings['collection_states']:
        return 'collection_status'
    values = settings['institutions']
    value = row.get('Institution / District', '')
    if settings['institution_match'] == 'exact_case_insensitive':
        value, values = value.casefold(), [v.strip().casefold() for v in values]
    if values and value not in values:
        return 'institution_filter'
    roll = row.get('Roll No', '')
    if settings['rolls'] and roll not in settings['rolls']:
        return 'roll_filter'
    if settings['roll_bounds']:
        lo, hi = settings['roll_bounds']
        if not roll.isascii() or not roll.isdigit() or not int(lo) <= int(roll) <= int(hi):
            return 'roll_bounds'
    if settings['row_ranges'] and not any(lo <= number <= hi for lo, hi in settings['row_ranges']):
        return 'row_range'
    if settings['result_statuses'] and row.get('Result Status', '') not in settings['result_statuses']:
        return 'result_status'
    return None
